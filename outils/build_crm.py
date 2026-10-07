# -*- coding: utf-8 -*-
"""Construit operations/crm-kpi-evenox.xlsx (CRM + dépenses + tableau de bord + import Google Ads)."""
import sys
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/user/Evenox/operations/crm-kpi-evenox.xlsx"

# ---------------------------------------------------------------- styles
NAVY = "1F3A5F"
F_HEAD = PatternFill("solid", fgColor=NAVY)
F_SUB = PatternFill("solid", fgColor="D9E2F3")
F_CALC = PatternFill("solid", fgColor="EDEDED")
F_INPUT = PatternFill("solid", fgColor="FFF9E6")
F_EX = PatternFill("solid", fgColor="FFE699")
F_TOTAL = PatternFill("solid", fgColor="DDEBF7")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color=NAVY)
SMALL_I = Font(italic=True, size=9, color="595959")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_L = Alignment(wrap_text=True, vertical="top", horizontal="left")
MONEY = '#,##0" $"'
MONEY2 = '#,##0.00" $"'
PCT = '0 %'
DT = 'yyyy-mm-dd hh:mm'
D = 'yyyy-mm-dd'
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
ORANGE_FILL = PatternFill("solid", fgColor="FFD966")
GREY_FILL = PatternFill("solid", fgColor="E7E6E6")
RED_FONT = Font(color="9C0006", bold=True)
GREEN_FONT = Font(color="006100", bold=True)

wb = Workbook()

# ---------------------------------------------------------------- reference lists (doc-exact)
CAMPAGNES = [  # libellé, $/jour (COMPTE-RENDU.md section 4)
    ("Google Search – Corporatif", 100),
    ("Meta – Leads", 80),
    ("Google Search – Mariage et privé", 60),
    ("ChatGPT Ads – Test", 30),
    ("Microsoft Ads", 20),
    ("Google Search – Marque", 10),
]
CAMP_NAMES = [c for c, _ in CAMPAGNES]
AUTRES_SOURCES = ["Fiche Google (organique)", "Référence / bouche-à-oreille",
                  "Appel direct (non attribué)", "Autre / non payé"]
TYPE_EVT = [  # code, libellé, segment
    ("corpo_party", "Party de bureau / des Fêtes", "Corporatif"),
    ("corpo_5a7", "5 à 7 ou team building", "Corporatif"),
    ("corpo_gala", "Gala, remise de prix ou lancement", "Corporatif"),
    ("muni_ecole", "Événement municipal ou scolaire", "Municipal / scolaire"),
    ("mariage", "Mariage", "Mariage"),
    ("prive", "Autre événement privé : anniversaire, shower…", "Privé"),
]
TYPE_CLIENT = [("entreprise", "Entreprise", 20), ("organisme", "Municipalité, école ou OBNL", 15),
               ("particulier", "Particulier ou couple", 5), ("agence", "Planificateur ou agence", 20)]
NB_INV = [("lt30", "Moins de 30", 0), ("30_74", "30 à 74", 6), ("75_149", "75 à 149", 12),
          ("150p", "150 et plus", 15)]
BUDGET = [("lt600", "Moins de 600 $", 0), ("600_999", "600 à 999 $", 10),
          ("1000_2499", "1 000 à 2 499 $", 15), ("2500_4999", "2 500 à 4 999 $", 30),
          ("5000p", "5 000 $ et plus", 40), ("inconnu", "Je ne sais pas encore", 10)]
ROLE = [("Je décide", 10), ("Je recommande, quelqu'un d'autre approuve", 6),
        ("Je compare des options pour un comité", 4)]
ECHEANCE = [("Cette semaine", 5), ("D'ici 30 jours", 3), ("Plus tard / je m'informe", 0)]
ROUTES = [("A", "Appel prioritaire"), ("B", "Soumission"), ("C", "Boutique"), ("D", "Hors zone")]
ETAPES = [  # étape, n°, probabilité (HubSpot, tracking-setup.md §9)
    ("Nouveau", "1", 0.10), ("Contacté", "2", 0.20), ("Qualifié", "3", 0.40),
    ("Proposition envoyée", "4", 0.50), ("Négociation", "5", 0.70),
    ("Dépôt payé (gagné)", "6", 1.0), ("Événement livré", "7", 1.0), ("Avis obtenu", "8", 1.0),
    ("Perdu", "—", 0.0), ("Disqualifié", "—", 0.0),
]
GAGNE = ["Dépôt payé (gagné)", "Événement livré", "Avis obtenu"]
RAISONS = ["Perdu – prix", "Perdu – concurrent", "Perdu – date", "Perdu – sans réponse",
           "Perdu – projet annulé", "Disqualifié – budget < seuil", "Disqualifié – hors zone",
           "Disqualifié – hors catalogue", "Disqualifié – pourriel"]
VILLES = ["Sainte-Thérèse", "Blainville", "Boisbriand", "Rosemère", "Lorraine", "Bois-des-Filion",
          "Mirabel", "Saint-Eustache", "Deux-Montagnes", "Terrebonne", "Saint-Jérôme", "Laval",
          "Montréal – centre-ville", "Montréal – autres arrondissements", "Longueuil / Rive-Sud",
          "Autre (plus de 40 km)"]
UTM_SOURCE = ["google", "bing", "meta", "chatgpt", "courriel", "sms", "gbp", "weddingwire", "yelp"]
UTM_MEDIUM = ["cpc", "paid_social", "email", "sms", "referral", "organic_local"]
UTM_CAMPAIGN = ["gads_corpo_fetes", "gads_corpo_5a7", "gads_mariage_deco", "gads_marque",
                "meta_corpo_fetes", "meta_mariage_deco", "meta_formulaire_instantane",
                "msft_corpo", "msft_mariage", "msft_marque",
                "chatgpt_corpo_fetes", "chatgpt_corpo_gala", "chatgpt_mariage"]
FORFAITS = ["5 à 7 (1 195 $)", "Party de bureau (1 995 $)", "Gala Signature (2 495 $)",
            "Décor WOW (899 $)", "Soirée Signature (1 449 $)", "Mariage Signature (1 899 $)",
            "Photobooth seul (dès 599 $)", "Sur mesure"]
OUI_NON = ["oui", "non"]

# ---------------------------------------------------------------- sheet: Listes (built first; referenced)
wsL = wb.active
wsL.title = "Listes"
list_ranges = {}


def put_table(col, title, headers, rows, key):
    c = col
    wsL.cell(1, c, title).font = BOLD
    for j, h in enumerate(headers):
        cell = wsL.cell(2, c + j, h)
        cell.fill, cell.font, cell.alignment = F_HEAD, WHITE_B, WRAP_C
    for i, r in enumerate(rows):
        r = r if isinstance(r, (tuple, list)) else (r,)
        for j, v in enumerate(r):
            cell = wsL.cell(3 + i, c + j, v)
            cell.border = BORDER
            if isinstance(v, float) and v <= 1:
                cell.number_format = PCT
    last = 2 + len(rows)
    list_ranges[key] = (c, last, len(headers))
    for j in range(len(headers)):
        wsL.column_dimensions[L(c + j)].width = 26 if j == 0 else 18
    return c + len(headers) + 1


col = 1
col = put_table(col, "Campagnes payées", ["campagne", "$/jour"], CAMPAGNES, "camp")
col = put_table(col, "Source / campagne (saisie manuelle)", ["campagne_manuelle"],
                CAMP_NAMES + AUTRES_SOURCES, "campman")
col = put_table(col, "type_evenement", ["code", "libellé", "segment"], TYPE_EVT, "type_evt")
col = put_table(col, "type_client", ["code", "points", "libellé"],
                [(a, p, b) for a, b, p in TYPE_CLIENT], "type_client")
col = put_table(col, "nb_invites", ["code", "points", "libellé"], [(a, p, b) for a, b, p in NB_INV], "nb")
col = put_table(col, "budget_declare", ["code", "points", "libellé"], [(a, p, b) for a, b, p in BUDGET], "budget")
col = put_table(col, "role_decision", ["libellé", "points"], ROLE, "role")
col = put_table(col, "echeance_decision", ["libellé", "points"], ECHEANCE, "ech")
col = put_table(col, "route", ["route", "libellé"], ROUTES, "route")
col = put_table(col, "etape", ["étape", "probabilité", "n°"], [(a, p, n) for a, n, p in ETAPES], "etape")
col = put_table(col, "raison_perte", ["raison"], RAISONS, "raison")
col = put_table(col, "ville", ["ville"], VILLES, "ville")
col = put_table(col, "utm_source", ["valeur"], UTM_SOURCE, "utm_source")
col = put_table(col, "utm_medium", ["valeur"], UTM_MEDIUM, "utm_medium")
col = put_table(col, "utm_campaign", ["valeur"], UTM_CAMPAIGN, "utm_campaign")
col = put_table(col, "forfait_propose", ["forfait"], FORFAITS, "forfait")
col = put_table(col, "oui / non", ["valeur"], OUI_NON, "ouinon")
wsL.freeze_panes = "A3"
wsL.sheet_properties.tabColor = "A6A6A6"


def lref(key, colofs=0, whole=False):
    c, last, n = list_ranges[key]
    if whole:
        return f"Listes!${L(c)}$3:${L(c + n - 1)}${last}"
    return f"Listes!${L(c + colofs)}$3:${L(c + colofs)}${last}"


# ---------------------------------------------------------------- sheet: Leads
ws = wb.create_sheet("Leads", 0)
FIRST, LAST = 3, 1002  # data rows
NROWS = LAST - FIRST + 1

# (key, libellé FR, largeur, nature) — les 39 premières colonnes = ordre exact de tracking-setup.md §9
COLS = [
    ("lead_id", "N° du lead", 14, "in"),
    ("date_heure_reception", "Date et heure de réception", 17, "dt"),
    ("date_heure_1er_contact", "Date et heure du 1er contact", 17, "dt"),
    ("delai_min", "Délai de 1er contact (min)", 11, "calc"),
    ("prenom", "Prénom", 12, "in"),
    ("nom", "Nom", 14, "in"),
    ("courriel", "Courriel", 26, "in"),
    ("telephone", "Téléphone", 14, "in"),
    ("entreprise", "Entreprise", 18, "in"),
    ("type_client", "Type de client", 13, "dv"),
    ("type_evenement", "Type d'événement", 14, "dv"),
    ("date_evenement", "Date de l'événement", 12, "d"),
    ("ville", "Ville", 20, "dv"),
    ("nb_invites", "Nombre d'invités", 10, "dv"),
    ("budget_declare", "Budget déclaré", 12, "dv"),
    ("route", "Route (formulaire)", 9, "dv"),
    ("lead_score", "Pointage (0-100)", 9, "calc"),
    ("etape", "Étape du pipeline", 20, "dv"),
    ("raison_perte", "Raison de perte / disqualification", 26, "dv"),
    ("forfait_propose", "Forfait proposé", 24, "dv"),
    ("montant_propose", "Montant proposé ($)", 12, "money"),
    ("montant_gagne", "Montant gagné ($)", 12, "money"),
    ("date_depot", "Date et heure du dépôt", 17, "dt"),
    ("utm_source", "utm_source", 11, "dv"),
    ("utm_medium", "utm_medium", 12, "dv"),
    ("utm_campaign", "utm_campaign", 20, "dv"),
    ("utm_content", "utm_content", 18, "in"),
    ("utm_term", "utm_term", 18, "in"),
    ("gclid", "gclid (Google)", 20, "in"),
    ("msclkid", "msclkid (Microsoft)", 18, "in"),
    ("fbclid", "fbclid (Meta)", 16, "in"),
    ("oppref", "oppref (ChatGPT)", 16, "in"),
    ("event_id", "event_id (déduplication)", 18, "in"),
    ("consent_marketing", "Consentement marketing", 11, "dv"),
    ("consent_mesure", "Consentement mesure", 11, "dv"),
    ("date_consentement", "Date du consentement", 17, "dt"),
    ("desabonne", "Désabonné", 10, "dv"),
    ("conseiller", "Conseiller", 14, "in"),
    ("notes", "Notes", 30, "in"),
    # --- champs complémentaires (formulaire-qualification.md) et colonnes calculées
    ("role_decision", "Rôle dans la décision", 22, "dv"),
    ("echeance_decision", "Échéance de décision", 16, "dv"),
    ("campagne_manuelle", "Source / campagne (forcer)", 24, "dv"),
    ("campagne", "Campagne (calculée)", 26, "calc"),
    ("segment", "Segment", 14, "calc"),
    ("route_calculee", "Route selon les règles", 10, "calc"),
    ("route_effective", "Route retenue", 9, "calc"),
    ("qualifie", "Lead qualifié (A ou B) 1/0", 10, "calc"),
    ("gagne", "Gagné (dépôt payé ou plus) 1/0", 10, "calc"),
    ("priorite_rappel", "Priorité de rappel", 22, "calc"),
    ("valeur_conv_provisoire", "Valeur de conversion provisoire ($)", 12, "calc"),
    ("valeur", "Valeur ($) : gagnée ou pondérée", 13, "calc"),
    ("est_exemple", "Ligne EXEMPLE 1/0", 9, "calc"),
    ("rang_import", "Rang import Google", 9, "calc"),
]
C = {k: L(i + 1) for i, (k, *_r) in enumerate(COLS)}
NDOC = 39


def rng(key, sheet="Leads"):
    return f"{sheet}!${C[key]}${FIRST}:${C[key]}${LAST}"


# title-ish row 1 (FR labels) and row 2 (technical names)
for i, (k, lab, w, kind) in enumerate(COLS, start=1):
    c1 = ws.cell(1, i, lab)
    c2 = ws.cell(2, i, k)
    c1.font, c1.alignment = WHITE_B, WRAP_C
    c1.fill = F_HEAD if i <= NDOC else PatternFill("solid", fgColor="2E75B6")
    if kind == "calc":
        c1.fill = PatternFill("solid", fgColor="595959")
    c2.font = Font(italic=True, size=9, color="1F3A5F")
    c2.fill, c2.alignment, c2.border = F_SUB, WRAP_C, BORDER
    ws.column_dimensions[L(i)].width = w
ws.row_dimensions[1].height = 48
ws.cell(1, 1).comment = Comment(
    "Ligne 1 : libellé. Ligne 2 : nom technique (identique à tracking-setup.md, section 9). "
    "Colonnes grises = formules, ne pas saisir.", "Évenox")


def f_row(r):
    g = lambda k: f"{C[k]}{r}"
    A = g("lead_id")
    out = {}
    out["delai_min"] = (f'=IF(OR({g("date_heure_reception")}="",{g("date_heure_1er_contact")}=""),"",'
                        f'ROUND(({g("date_heure_1er_contact")}-{g("date_heure_reception")})*1440,0))')
    bud, tc, nb, de = g("budget_declare"), g("type_client"), g("nb_invites"), g("date_evenement")
    rec, te, vi = g("date_heure_reception"), g("type_evenement"), g("ville")
    out["lead_score"] = (
        f'=IF({A}="","",'
        f'IFERROR(VLOOKUP({bud},{lref("budget", whole=True)},2,FALSE),0)'
        f'+IFERROR(VLOOKUP({tc},{lref("type_client", whole=True)},2,FALSE),0)'
        f'+IFERROR(VLOOKUP({nb},{lref("nb", whole=True)},2,FALSE),0)'
        f'+IF(OR({de}="",{rec}=""),0,IF({de}-INT({rec})<10,8,IF({de}-INT({rec})<=90,10,5)))'
        f'+IFERROR(VLOOKUP({g("role_decision")},{lref("role", whole=True)},2,FALSE),0)'
        f'+IFERROR(VLOOKUP({g("echeance_decision")},{lref("ech", whole=True)},2,FALSE),0))')
    uc, us = g("utm_campaign"), g("utm_source")
    out["campagne"] = (
        f'=IF({A}="","",IF({g("campagne_manuelle")}<>"",{g("campagne_manuelle")},'
        f'IF(LEFT({uc},11)="gads_marque","{CAMP_NAMES[5]}",'
        f'IF(LEFT({uc},12)="gads_mariage","{CAMP_NAMES[2]}",'
        f'IF(LEFT({uc},10)="gads_corpo","{CAMP_NAMES[0]}",'
        f'IF(LEFT({uc},5)="meta_","{CAMP_NAMES[1]}",'
        f'IF(LEFT({uc},5)="msft_","{CAMP_NAMES[4]}",'
        f'IF(LEFT({uc},8)="chatgpt_","{CAMP_NAMES[3]}",'
        f'IF({us}="gbp","Fiche Google (organique)","Autre / non payé")))))))))')
    out["segment"] = f'=IF({te}="","",IFERROR(VLOOKUP({te},{lref("type_evt", whole=True)},3,FALSE),""))'
    out["route_calculee"] = (
        f'=IF({bud}="","",'
        f'IF(AND({vi}="Autre (plus de 40 km)",OR({bud}="lt600",{bud}="600_999",{bud}="1000_2499")),"D",'
        f'IF(OR({bud}="lt600",AND({bud}="600_999",{te}<>"mariage",{te}<>"prive")),"C",'
        f'IF(OR({bud}="2500_4999",{bud}="5000p",{tc}="entreprise",{tc}="organisme",{tc}="agence"),"A",'
        f'IF(OR({bud}="1000_2499",{bud}="600_999"),"B",'
        f'IF({bud}="inconnu",IF(OR({nb}="75_149",{nb}="150p"),"A","B"),""))))))')
    out["route_effective"] = f'=IF({A}="","",IF({g("route")}<>"",{g("route")},{g("route_calculee")}))'
    re_ = g("route_effective")
    out["qualifie"] = f'=IF({A}="","",IF(OR({re_}="A",{re_}="B"),1,0))'
    et = g("etape")
    out["gagne"] = f'=IF({A}="","",IF(OR(' + ",".join(f'{et}="{s}"' for s in GAGNE) + '),1,0))'
    sc = g("lead_score")
    out["priorite_rappel"] = (f'=IF({A}="","",IF({sc}>70,"Immédiat – propriétaire",'
                              f'IF({sc}>=40,"15 min – personne de garde","Selon la route")))')
    out["valeur_conv_provisoire"] = f'=IF({A}="","",IF({re_}="A",400,IF({re_}="B",150,0)))'
    mg, mp = g("montant_gagne"), g("montant_propose")
    out["valeur"] = (f'=IF({A}="","",IF({g("gagne")}=1,N({mg}),'
                     f'IF(N({mp})>0,ROUND(N({mp})*IFERROR(VLOOKUP({et},{lref("etape", whole=True)},2,FALSE),0),0),0)))')
    out["est_exemple"] = f'=IF(LEFT({A},7)="EXEMPLE",1,0)'
    gc, dd, ex, ga = g("gclid"), g("date_depot"), g("est_exemple"), g("gagne")
    rr = lambda k: f"${C[k]}${FIRST}:{C[k]}{r}"
    out["rang_import"] = (
        f'=IF(AND({ga}=1,{ex}=0,{gc}<>"",{dd}<>"",N({mg})>0),'
        f'COUNTIFS({rr("gagne")},1,{rr("est_exemple")},0,{rr("gclid")},"<>",'
        f'{rr("date_depot")},"<>",{rr("montant_gagne")},">0"),"")')
    return out


fmt_by_kind = {"dt": DT, "d": D, "money": MONEY}
for r in range(FIRST, LAST + 1):
    fr = f_row(r)
    for i, (k, lab, w, kind) in enumerate(COLS, start=1):
        cell = ws.cell(r, i)
        if k in fr:
            cell.value = fr[k]
            cell.fill = F_CALC
        if kind in fmt_by_kind:
            cell.number_format = fmt_by_kind[kind]
        if k in ("valeur_conv_provisoire", "valeur"):
            cell.number_format = MONEY
        cell.border = BORDER

# --- 5 EXEMPLE rows
EX = [
    dict(lead_id="EXEMPLE-001", date_heure_reception=datetime(2026, 10, 1, 10, 12),
         date_heure_1er_contact=datetime(2026, 10, 1, 10, 16), prenom="Exemple", nom="Corpo-Party",
         courriel="exemple1@exemple.com", telephone="+15145550101", entreprise="Entreprise Exemple inc.",
         type_client="entreprise", type_evenement="corpo_party", date_evenement=date(2026, 12, 11),
         ville="Laval", nb_invites="75_149", budget_declare="2500_4999", route="A",
         etape="Dépôt payé (gagné)", forfait_propose="Gala Signature (2 495 $)", montant_propose=2495,
         montant_gagne=2495, date_depot=datetime(2026, 10, 3, 14, 30), utm_source="google",
         utm_medium="cpc", utm_campaign="gads_corpo_fetes", utm_content="rsa_prixfixe_v1",
         utm_term="party de bureau", gclid="EXEMPLE-GCLID-001", event_id="ex-0001",
         consent_marketing="oui", consent_mesure="oui", date_consentement=datetime(2026, 10, 1, 10, 12),
         desabonne="non", conseiller="Alexandre", notes="EXEMPLE : supprimer cette ligne avant l'usage réel.",
         role_decision="Je décide", echeance_decision="Cette semaine"),
    dict(lead_id="EXEMPLE-002", date_heure_reception=datetime(2026, 10, 2, 19, 40),
         date_heure_1er_contact=datetime(2026, 10, 2, 20, 2), prenom="Exemple", nom="Mariage",
         courriel="exemple2@exemple.com", telephone="+14505550102", type_client="particulier",
         type_evenement="mariage", date_evenement=date(2027, 7, 17), ville="Mirabel",
         nb_invites="75_149", budget_declare="1000_2499", route="B", etape="Proposition envoyée",
         forfait_propose="Soirée Signature (1 449 $)", montant_propose=1449, utm_source="google",
         utm_medium="cpc", utm_campaign="gads_mariage_deco", utm_content="rsa_dates_v1",
         utm_term="decor mariage", gclid="EXEMPLE-GCLID-002", event_id="ex-0002",
         consent_marketing="non", consent_mesure="oui", date_consentement=datetime(2026, 10, 2, 19, 40),
         desabonne="non", conseiller="Alexandre",
         notes="EXEMPLE : délai de 22 min sur une route B, la ligne est signalée en rouge.",
         role_decision="Je décide", echeance_decision="D'ici 30 jours"),
    dict(lead_id="EXEMPLE-003", date_heure_reception=datetime(2026, 10, 3, 11, 5),
         prenom="Exemple", nom="Privé", courriel="exemple3@exemple.com", telephone="+15145550103",
         type_client="particulier", type_evenement="prive", date_evenement=date(2026, 11, 7),
         ville="Boisbriand", nb_invites="lt30", budget_declare="lt600", route="C", etape="Disqualifié",
         raison_perte="Disqualifié – budget < seuil", utm_source="meta", utm_medium="paid_social",
         utm_campaign="meta_mariage_deco", utm_content="reel_avantapres_v2", fbclid="EXEMPLE-FBCLID-003",
         event_id="ex-0003", consent_marketing="non", consent_mesure="non", desabonne="non",
         conseiller="Alexandre", notes="EXEMPLE : route C, redirigé vers la boutique.",
         role_decision="Je décide", echeance_decision="Cette semaine"),
    dict(lead_id="EXEMPLE-004", date_heure_reception=datetime(2026, 10, 5, 9, 30),
         date_heure_1er_contact=datetime(2026, 10, 5, 9, 33), prenom="Exemple", nom="Gala",
         courriel="exemple4@exemple.com", telephone="+14505550104", entreprise="Ville Exemple",
         type_client="organisme", type_evenement="corpo_gala", date_evenement=date(2027, 2, 20),
         ville="Terrebonne", nb_invites="150p", budget_declare="5000p", route="A", etape="Négociation",
         forfait_propose="Sur mesure", montant_propose=3200, utm_source="chatgpt", utm_medium="cpc",
         utm_campaign="chatgpt_corpo_fetes", utm_content="card_dates_dec_v1", oppref="EXEMPLE-OPPREF-004",
         event_id="ex-0004", consent_marketing="oui", consent_mesure="non",
         date_consentement=datetime(2026, 10, 5, 9, 30), desabonne="non", conseiller="Alexandre",
         notes="EXEMPLE : comité, relance la veille de la réunion.",
         role_decision="Je compare des options pour un comité", echeance_decision="D'ici 30 jours"),
    dict(lead_id="EXEMPLE-005", date_heure_reception=datetime(2026, 10, 6, 13, 50),
         date_heure_1er_contact=datetime(2026, 10, 6, 13, 58), prenom="Exemple", nom="5à7",
         courriel="exemple5@exemple.com", telephone="+15145550105", entreprise="PME Exemple",
         type_client="entreprise", type_evenement="corpo_5a7", date_evenement=date(2026, 11, 26),
         ville="Montréal – centre-ville", nb_invites="30_74", budget_declare="1000_2499", route="A",
         etape="Perdu", raison_perte="Perdu – concurrent", forfait_propose="5 à 7 (1 195 $)",
         montant_propose=1195, utm_source="bing", utm_medium="cpc", utm_campaign="msft_corpo",
         utm_content="rsa_prixfixe_v1", utm_term="5 a 7 entreprise", msclkid="EXEMPLE-MSCLKID-005",
         event_id="ex-0005", consent_marketing="non", consent_mesure="non", desabonne="non",
         conseiller="Alexandre", notes="EXEMPLE : concurrent = [nom] ; noter lequel.",
         role_decision="Je recommande, quelqu'un d'autre approuve", echeance_decision="Cette semaine"),
]
for idx, row in enumerate(EX):
    r = FIRST + idx
    for k, v in row.items():
        ws[f"{C[k]}{r}"] = v
    for i, (k, *_x) in enumerate(COLS, start=1):
        if not str(ws.cell(r, i).value or "").startswith("="):
            ws.cell(r, i).fill = F_EX
    ws[f"{C['lead_id']}{r}"].font = Font(bold=True, color="C00000")

# --- data validation
def add_dv(key, formula, allow_other=False, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if allow_other:
        dv.errorStyle = "warning"
        dv.error = "Valeur hors liste : vérifiez qu'elle respecte les conventions (tracking-setup.md)."
        dv.errorTitle = "Valeur non standard"
    else:
        dv.error = "Choisissez une valeur dans la liste."
        dv.errorTitle = "Valeur non permise"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt, dv.promptTitle, dv.showInputMessage = prompt, "Aide", True
    ws.add_data_validation(dv)
    dv.add(f"{C[key]}{FIRST}:{C[key]}{LAST}")


add_dv("type_client", lref("type_client"))
add_dv("type_evenement", lref("type_evt"), prompt="Code du formulaire : voir la feuille Listes pour le libellé.")
add_dv("ville", lref("ville"))
add_dv("nb_invites", lref("nb"))
add_dv("budget_declare", lref("budget"))
add_dv("route", lref("route"), prompt="A, B, C ou D, telle qu'envoyée par le formulaire. Vide = route calculée.")
add_dv("etape", lref("etape"))
add_dv("raison_perte", lref("raison"), prompt="Obligatoire si l'étape est Perdu ou Disqualifié.")
add_dv("forfait_propose", lref("forfait"), allow_other=True)
add_dv("utm_source", lref("utm_source"), allow_other=True)
add_dv("utm_medium", lref("utm_medium"), allow_other=True)
add_dv("utm_campaign", lref("utm_campaign"), allow_other=True)
add_dv("consent_marketing", lref("ouinon"))
add_dv("consent_mesure", lref("ouinon"))
add_dv("desabonne", lref("ouinon"))
add_dv("role_decision", lref("role"))
add_dv("echeance_decision", lref("ech"))
add_dv("campagne_manuelle", lref("campman"),
       prompt="Laisser vide : la campagne est déduite de utm_campaign. Remplir pour un appel, une référence, etc.")
for k in ("date_heure_reception", "date_heure_1er_contact", "date_depot", "date_consentement", "date_evenement"):
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2024,1,1)", allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "Entrez une date (AAAA-MM-JJ HH:MM).", "Date invalide", True
    ws.add_data_validation(dv)
    dv.add(f"{C[k]}{FIRST}:{C[k]}{LAST}")
for k in ("montant_propose", "montant_gagne"):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.error, dv.showErrorMessage = "Montant en dollars, sans taxes.", True
    ws.add_data_validation(dv)
    dv.add(f"{C[k]}{FIRST}:{C[k]}{LAST}")

# --- conditional formatting (Leads)
data_area = f"A{FIRST}:{L(len(COLS))}{LAST}"
ws.conditional_formatting.add(
    f"A{FIRST}:{L(NDOC)}{LAST}",
    FormulaRule(formula=[f'AND(ISNUMBER(${C["delai_min"]}{FIRST}),${C["delai_min"]}{FIRST}>15,'
                         f'OR(${C["route_effective"]}{FIRST}="A",${C["route_effective"]}{FIRST}="B"))'],
                fill=RED_FILL, stopIfTrue=False))
ws.conditional_formatting.add(
    f"{C['route']}{FIRST}:{C['route']}{LAST}",
    FormulaRule(formula=[f'AND({C["route"]}{FIRST}<>"",{C["route_calculee"]}{FIRST}<>"",'
                         f'{C["route"]}{FIRST}<>{C["route_calculee"]}{FIRST})'],
                fill=ORANGE_FILL))
ws.conditional_formatting.add(
    f"{C['raison_perte']}{FIRST}:{C['raison_perte']}{LAST}",
    FormulaRule(formula=[f'AND(OR({C["etape"]}{FIRST}="Perdu",{C["etape"]}{FIRST}="Disqualifié"),'
                         f'{C["raison_perte"]}{FIRST}="")'], fill=RED_FILL))
ws.conditional_formatting.add(
    f"{C['lead_score']}{FIRST}:{C['lead_score']}{LAST}",
    FormulaRule(formula=[f'AND(ISNUMBER({C["lead_score"]}{FIRST}),{C["lead_score"]}{FIRST}>70)'],
                fill=GREEN_FILL, font=GREEN_FONT))
ws.conditional_formatting.add(
    f"{C['etape']}{FIRST}:{C['etape']}{LAST}",
    FormulaRule(formula=[f'{C["gagne"]}{FIRST}=1'], fill=GREEN_FILL, font=GREEN_FONT))

ws.freeze_panes = "B3"
ws.auto_filter.ref = f"A2:{L(len(COLS))}{LAST}"
ws.sheet_properties.tabColor = NAVY

# ---------------------------------------------------------------- sheet: Dépenses
wd = wb.create_sheet("Dépenses", 1)
DS = "'Dépenses'"
wd["A1"] = "Dépenses publicitaires quotidiennes par campagne ($ CA, avant taxes)"
wd["A1"].font = TITLE
wd["A2"] = "Saisir chaque matin la dépense de la veille (rapport de chaque plateforme). Cellules jaunes = saisie."
wd["A2"].font = SMALL_I
hdr_row, D_FIRST = 7, 8
start_d, end_d = date(2026, 10, 1), date(2027, 12, 31)
ndays = (end_d - start_d).days + 1
D_LAST = D_FIRST + ndays - 1
camp_cols = {name: L(3 + i) for i, name in enumerate(CAMP_NAMES)}  # C..H
TOT, BUD, ECA, ECP, NOTE = "I", "J", "K", "L", "M"

labels = [(3, "Budget cible ($/jour)"), (4, "Dépensé (cumul)"), (5, "Budget des jours saisis"), (6, "Écart cumulé ($)")]
for r, lab in labels:
    wd[f"B{r}"] = lab
    wd[f"B{r}"].font = BOLD
    wd[f"B{r}"].border = BORDER
for name, budget_ in CAMPAGNES:
    c = camp_cols[name]
    wd[f"{c}3"] = budget_
    wd[f"{c}3"].fill = F_INPUT
    wd[f"{c}4"] = f"=SUM({c}{D_FIRST}:{c}{D_LAST})"
    wd[f"{c}5"] = f"={c}3*COUNT({c}{D_FIRST}:{c}{D_LAST})"
    wd[f"{c}6"] = f"={c}4-{c}5"
    for r in (3, 4, 5, 6):
        wd[f"{c}{r}"].number_format = MONEY
        wd[f"{c}{r}"].border = BORDER
wd[f"{TOT}3"] = f"=SUM(C3:H3)"
wd[f"{TOT}4"] = f"=SUM(C4:H4)"
wd[f"{TOT}5"] = f"=SUM(C5:H5)"
wd[f"{TOT}6"] = f"={TOT}4-{TOT}5"
for r in (3, 4, 5, 6):
    wd[f"{TOT}{r}"].number_format = MONEY
    wd[f"{TOT}{r}"].font = BOLD
    wd[f"{TOT}{r}"].border = BORDER
wd[f"{BUD}3"] = "← total visé : 300 $/jour. Ajuster lors de la bascule saisonnière (COMPTE-RENDU.md)."
wd[f"{BUD}3"].font = SMALL_I

heads = ["Date", "Jour"] + CAMP_NAMES + ["Total réel", "Budget cible du jour", "Écart ($)", "Écart (%)", "Notes"]
for i, h in enumerate(heads, start=1):
    c = wd.cell(hdr_row, i, h)
    c.fill, c.font, c.alignment, c.border = F_HEAD, WHITE_B, WRAP_C, BORDER
wd.row_dimensions[hdr_row].height = 45
JOURS = '"lundi","mardi","mercredi","jeudi","vendredi","samedi","dimanche"'
for n in range(ndays):
    r = D_FIRST + n
    wd[f"A{r}"] = start_d + timedelta(days=n)
    wd[f"A{r}"].number_format = D
    wd[f"B{r}"] = f"=CHOOSE(WEEKDAY(A{r},2),{JOURS})"
    for name in CAMP_NAMES:
        c = camp_cols[name]
        wd[f"{c}{r}"].number_format = MONEY2
        wd[f"{c}{r}"].fill = F_INPUT
    wd[f"{TOT}{r}"] = f'=IF(COUNT(C{r}:H{r})=0,"",SUM(C{r}:H{r}))'
    wd[f"{BUD}{r}"] = f"=${TOT}$3"
    wd[f"{ECA}{r}"] = f'=IF({TOT}{r}="","",{TOT}{r}-{BUD}{r})'
    wd[f"{ECP}{r}"] = f'=IF(OR({TOT}{r}="",N({BUD}{r})=0),"",{ECA}{r}/{BUD}{r})'
    wd[f"{TOT}{r}"].number_format = MONEY2
    wd[f"{BUD}{r}"].number_format = MONEY
    wd[f"{ECA}{r}"].number_format = MONEY2
    wd[f"{ECP}{r}"].number_format = PCT
    for col_ in "ABCDEFGHIJKLM":
        wd[f"{col_}{r}"].border = BORDER
    for col_ in ("B", TOT, BUD, ECA, ECP):
        wd[f"{col_}{r}"].fill = F_CALC
    if (start_d + timedelta(days=n)).weekday() == 0:
        for col_ in "AB":
            wd[f"{col_}{r}"].font = BOLD
wd.conditional_formatting.add(f"{ECP}{D_FIRST}:{ECP}{D_LAST}",
                              FormulaRule(formula=[f'AND(ISNUMBER({ECP}{D_FIRST}),{ECP}{D_FIRST}>0.1)'],
                                          fill=RED_FILL))
wd.conditional_formatting.add(f"{ECP}{D_FIRST}:{ECP}{D_LAST}",
                              FormulaRule(formula=[f'AND(ISNUMBER({ECP}{D_FIRST}),{ECP}{D_FIRST}<-0.1)'],
                                          fill=ORANGE_FILL))
for name in CAMP_NAMES:
    c = camp_cols[name]
    wd.conditional_formatting.add(f"{c}6", FormulaRule(formula=[f"{c}6>0"], font=RED_FONT))
dvn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dvn.error, dvn.showErrorMessage = "Dépense en dollars (nombre positif).", True
wd.add_data_validation(dvn)
dvn.add(f"C{D_FIRST}:H{D_LAST}")
widths = {"A": 12, "B": 22, TOT: 12, BUD: 12, ECA: 11, ECP: 10, NOTE: 30}
for k, v in widths.items():
    wd.column_dimensions[k].width = v
for c in camp_cols.values():
    wd.column_dimensions[c].width = 15
wd.freeze_panes = f"C{D_FIRST}"
wd.sheet_properties.tabColor = "C55A11"

# ---------------------------------------------------------------- sheet: Tableau de bord
wt = wb.create_sheet("Tableau de bord", 2)
wt["B1"] = "Tableau de bord des indicateurs — Évenox (revue chaque lundi, 9 h)"
wt["B1"].font = TITLE
wt["B2"] = ("Leads comptés par date de réception (cohorte) ; contrats, revenu et ROAS = leads de la même cohorte "
            "qui ont atteint « Dépôt payé (gagné) » ou plus. Lead qualifié (LQ) = route A ou B.")
wt["B2"].font = SMALL_I

params = [
    (4, "Semaine débutant le (lundi)", "=TODAY()-WEEKDAY(TODAY(),2)+1", D),
    (5, "Mois analysé (1er jour du mois)", "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", D),
    (6, "Inclure les lignes EXEMPLE?", "Non", None),
    (7, "(critère interne, ne pas modifier)", '=IF(C6="Oui",1,0)', None),
]
for r, lab, val, nf in params:
    wt[f"B{r}"] = lab
    wt[f"B{r}"].font = BOLD if r < 7 else SMALL_I
    wt[f"C{r}"] = val
    wt[f"C{r}"].border = BORDER
    if r < 7:
        wt[f"C{r}"].fill = F_INPUT
    if nf:
        wt[f"C{r}"].number_format = nf
wt["D4"] = "Formule par défaut = lundi de la semaine en cours ; inscrire une date pour une autre semaine."
wt["D5"] = "Formule par défaut = mois en cours ; inscrire le 1er d'un mois (ex. 2026-11-01)."
for c in ("D4", "D5"):
    wt[c].font = SMALL_I
dv_on = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False)
wt.add_data_validation(dv_on)
dv_on.add("C6")
W_S, W_E, M_S, M_E, EXC = "$C$4", "($C$4+7)", "$C$5", "EDATE($C$5,1)", "$C$7"
W_PS = "($C$4-7)"

# targets
targets = [
    ("cplq", "Coût par lead qualifié (max)", 150, MONEY),
    ("pctq", "% de leads qualifiés (min)", 0.40, PCT),
    ("sig", "Taux de signature des LQ (min)", 0.25, PCT),
    ("panier", "Panier moyen (min)", 1500, MONEY),
    ("roas", "ROAS (min)", 2.5, "0.0"),
    ("reduire", "Règle « Réduire 30 % » : coût par LQ >", "=2*{cplq}", MONEY),
    ("depmin", "Dépense minimale avant décision (par campagne)", 300, MONEY),
    ("augm", "Règle « Augmenter +20 % » : coût par LQ <", 100, MONEY),
    ("gpt_lq", "ChatGPT au jour 30 : LQ minimum", 2, "0"),
    ("gpt_dep", "ChatGPT : dépense de référence (≈ 30 jours × 30 $)", 900, MONEY),
    ("delai", "Délai de réponse médian visé (min), routes A et B", 5, "0"),
]
wt["F3"] = "Cibles et seuils (tracking-setup.md §10, COMPTE-RENDU.md §4)"
wt["F3"].font = BOLD
T = {}
for i, (k, lab, v, nf) in enumerate(targets):
    r = 4 + i
    T[k] = f"$I${r}"
for i, (k, lab, v, nf) in enumerate(targets):
    r = 4 + i
    wt[f"F{r}"] = lab
    wt.merge_cells(f"F{r}:H{r}")
    wt[f"I{r}"] = v.format(**T) if isinstance(v, str) else v
    wt[f"I{r}"].number_format = nf
    wt[f"I{r}"].fill = F_INPUT
    wt[f"I{r}"].border = BORDER
    wt[f"F{r}"].border = BORDER

BLOCK_HEAD = ["Campagne", "Dépenses", "Budget cible (jours saisis)", "Écart vs budget", "Leads (tous)", "Leads qualifiés (LQ)",
              "% qualifiés", "Coût par lead", "Coût par LQ", "Contrats (dépôts)", "Taux de signature",
              "Revenu (montant gagné)", "Panier moyen", "ROAS", "Pipeline pondéré", "Décision"]
BCOLS = [L(2 + i) for i in range(len(BLOCK_HEAD))]  # B..Q
(cB, cDep, cBud, cEca, cLeads, cLQ, cPct, cCPL, cCPLQ, cCtr, cSig, cRev, cPan, cRoas, cPipe, cDec) = BCOLS


def crit(start, end, extra=""):
    return (f'{rng("date_heure_reception")},">="&{start},{rng("date_heure_reception")},"<"&{end},'
            f'{rng("est_exemple")},"<="&{EXC}{extra}')


def build_block(top, title, start, end, days_expr, weekly):
    wt[f"B{top}"] = title
    wt[f"B{top}"].font = Font(bold=True, size=12, color=NAVY)
    h = top + 1
    heads = BLOCK_HEAD + (["Sem. préc. : dépenses", "Sem. préc. : LQ"] if weekly else [])
    for i, txt in enumerate(heads):
        c = wt.cell(h, 2 + i, txt)
        c.fill, c.font, c.alignment, c.border = F_HEAD, WHITE_B, WRAP_C, BORDER
    wt.row_dimensions[h].height = 42
    rows = {}
    r = h + 1
    first = r
    for name in CAMP_NAMES + ["Autres sources (non payées)", "Total"]:
        rows[name] = r
        r += 1
    last_c = rows[CAMP_NAMES[-1]]
    r_aut, r_tot = rows["Autres sources (non payées)"], rows["Total"]
    for name, r in rows.items():
        wt[f"{cB}{r}"] = name
        is_c = name in CAMP_NAMES
        camp_ref = f'{rng("campagne")},${cB}{r},'
        if is_c:
            dcol = camp_cols[name]
            dep = f"SUMIFS({DS}!${dcol}${D_FIRST}:${dcol}${D_LAST},{DS}!$A${D_FIRST}:$A${D_LAST},\">=\"&{start},{DS}!$A${D_FIRST}:$A${D_LAST},\"<\"&{end})"
            wt[f"{cDep}{r}"] = "=" + dep
            wt[f"{cBud}{r}"] = (f"={DS}!${dcol}$3*COUNTIFS({DS}!$A${D_FIRST}:$A${D_LAST},\">=\"&{start},"
                                  f"{DS}!$A${D_FIRST}:$A${D_LAST},\"<\"&{end},{DS}!${dcol}${D_FIRST}:${dcol}${D_LAST},\"<>\")")
            wt[f"{cLeads}{r}"] = f"=COUNTIFS({camp_ref}{crit(start, end)})"
            wt[f"{cLQ}{r}"] = f"=COUNTIFS({camp_ref}{crit(start, end, ',' + rng('qualifie') + ',1')})"
            wt[f"{cCtr}{r}"] = f"=COUNTIFS({camp_ref}{crit(start, end, ',' + rng('gagne') + ',1')})"
            wt[f"{cRev}{r}"] = f"=SUMIFS({rng('montant_gagne')},{camp_ref}{crit(start, end, ',' + rng('gagne') + ',1')})"
            wt[f"{cPipe}{r}"] = f"=SUMIFS({rng('valeur')},{camp_ref}{crit(start, end)})"
            if weekly:
                pdep = (f"SUMIFS({DS}!${dcol}${D_FIRST}:${dcol}${D_LAST},{DS}!$A${D_FIRST}:$A${D_LAST},"
                        f"\">=\"&{W_PS},{DS}!$A${D_FIRST}:$A${D_LAST},\"<\"&{W_S})")
                wt[f"R{r}"] = "=" + pdep
                wt[f"S{r}"] = f"=COUNTIFS({camp_ref}{crit(W_PS, W_S, ',' + rng('qualifie') + ',1')})"
        elif name.startswith("Autres"):
            wt[f"{cDep}{r}"] = 0
            wt[f"{cBud}{r}"] = 0
            for col_, extra in ((cLeads, ""), (cLQ, "," + rng("qualifie") + ",1"), (cCtr, "," + rng("gagne") + ",1")):
                wt[f"{col_}{r}"] = f"=COUNTIFS({crit(start, end, extra)})-SUM({col_}{first}:{col_}{last_c})"
            wt[f"{cRev}{r}"] = (f"=SUMIFS({rng('montant_gagne')},{crit(start, end, ',' + rng('gagne') + ',1')})"
                                f"-SUM({cRev}{first}:{cRev}{last_c})")
            wt[f"{cPipe}{r}"] = f"=SUMIFS({rng('valeur')},{crit(start, end)})-SUM({cPipe}{first}:{cPipe}{last_c})"
            if weekly:
                wt[f"R{r}"] = 0
                wt[f"S{r}"] = (f"=COUNTIFS({crit(W_PS, W_S, ',' + rng('qualifie') + ',1')})"
                               f"-SUM(S{first}:S{last_c})")
        else:  # total
            for col_ in (cDep, cBud, cLeads, cLQ, cCtr, cRev, cPipe) + (("R", "S") if weekly else ()):
                wt[f"{col_}{r}"] = f"=SUM({col_}{first}:{col_}{r_aut})"
        # ratios
        wt[f"{cEca}{r}"] = f"={cDep}{r}-{cBud}{r}"
        wt[f"{cPct}{r}"] = f'=IF({cLeads}{r}=0,"—",{cLQ}{r}/{cLeads}{r})'
        wt[f"{cCPL}{r}"] = f'=IF(OR({cLeads}{r}=0,{cDep}{r}=0),"—",{cDep}{r}/{cLeads}{r})'
        wt[f"{cCPLQ}{r}"] = f'=IF(OR({cLQ}{r}=0,{cDep}{r}=0),"—",{cDep}{r}/{cLQ}{r})'
        wt[f"{cSig}{r}"] = f'=IF({cLQ}{r}=0,"—",{cCtr}{r}/{cLQ}{r})'
        wt[f"{cPan}{r}"] = f'=IF({cCtr}{r}=0,"—",{cRev}{r}/{cCtr}{r})'
        wt[f"{cRoas}{r}"] = f'=IF({cDep}{r}=0,"—",{cRev}{r}/{cDep}{r})'
        # decision
        dep, lq, ctr = f"{cDep}{r}", f"{cLQ}{r}", f"{cCtr}{r}"
        cplq_n = f"IFERROR({dep}/{lq},1E+99)"
        sig_n = f"IFERROR({ctr}/{lq},0)"
        if is_c and weekly:
            prev_bad = f"AND(R{r}>0,IFERROR(R{r}/S{r},1E+99)>{T['reduire']})"
            cur_bad = f"AND({dep}>0,{cplq_n}>{T['reduire']})"
            wt[f"{cDec}{r}"] = (
                f'=IF(AND({dep}+R{r}>={T["depmin"]},{cur_bad},{prev_bad}),"Réduire 30 %",'
                f'IF(AND({dep}>0,{lq}>0,{cplq_n}<{T["augm"]},{sig_n}>={T["sig"]}),"Augmenter +20 %",'
                f'IF({dep}+R{r}<{T["depmin"]},"Garder (données insuffisantes)","Garder")))')
        elif is_c:
            couper = ""
            if name == "ChatGPT Ads – Test":
                couper = (f'IF(AND({dep}>={T["gpt_dep"]},OR({lq}<{T["gpt_lq"]},{cplq_n}>2*$C${top - 1})),'
                          f'"Couper",')
            wt[f"{cDec}{r}"] = (
                f'=IF({dep}<{T["depmin"]},"Garder (données insuffisantes)",{couper}'
                f'IF({cplq_n}>{T["reduire"]},"Réduire 30 %",'
                f'IF(AND({lq}>0,{cplq_n}<{T["augm"]},{sig_n}>={T["sig"]}),"Augmenter +20 %","Garder")))'
                + (")" if couper else ""))
        else:
            wt[f"{cDec}{r}"] = "—"
        # formats
        for col_ in (cDep, cBud, cEca, cCPL, cCPLQ, cRev, cPan, cPipe) + (("R",) if weekly else ()):
            wt[f"{col_}{r}"].number_format = MONEY
        for col_ in (cPct, cSig):
            wt[f"{col_}{r}"].number_format = PCT
        wt[f"{cRoas}{r}"].number_format = "0.00"
        for i in range(len(heads)):
            cell = wt.cell(r, 2 + i)
            cell.border = BORDER
            if name == "Total":
                cell.fill, cell.font = F_TOTAL, BOLD
            elif name.startswith("Autres"):
                cell.font = Font(italic=True, color="595959")
            if i > 0:
                cell.alignment = Alignment(horizontal="right")
        wt[f"{cDec}{r}"].alignment = Alignment(horizontal="center")
    # conditional formatting vs targets
    a, b = first, r_tot
    def cf(col_, good, bad):
        rg = f"{col_}{a}:{col_}{b}"
        wt.conditional_formatting.add(rg, FormulaRule(formula=[f"AND(ISNUMBER({col_}{a}),{good.format(x=f'{col_}{a}')})"],
                                                      fill=GREEN_FILL, font=GREEN_FONT))
        wt.conditional_formatting.add(rg, FormulaRule(formula=[f"AND(ISNUMBER({col_}{a}),{bad.format(x=f'{col_}{a}')})"],
                                                      fill=RED_FILL, font=RED_FONT))
    cf(cCPLQ, "{x}<=" + T["cplq"], "{x}>" + T["cplq"])
    cf(cPct, "{x}>=" + T["pctq"], "{x}<" + T["pctq"])
    cf(cSig, "{x}>=" + T["sig"], "{x}<" + T["sig"])
    cf(cPan, "{x}>=" + T["panier"], "{x}<" + T["panier"])
    cf(cRoas, "{x}>=" + T["roas"], "{x}<" + T["roas"])
    rg = f"{cDec}{a}:{cDec}{b}"
    for prefix, fill, font in (("Couper", PatternFill("solid", fgColor="C00000"), Font(bold=True, color="FFFFFF")),
                               ("Réduire", RED_FILL, RED_FONT),
                               ("Augmenter", GREEN_FILL, GREEN_FONT),
                               ("Garder", GREY_FILL, BOLD)):
        wt.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({cDec}{a},{len(prefix)})="{prefix}"'],
                                                      fill=fill, font=font))
    return rows


# weekly
WT = 17
wt[f"B{WT - 1}"] = None
w_rows = build_block(WT, "Semaine : du lundi indiqué en C4 au dimanche suivant", W_S, W_E, "7", True)
wt[f"B{WT - 1}"] = ("Décision hebdo (tracking-setup.md §10) : « Réduire 30 % » si coût par LQ > 2 × la cible "
                    "2 semaines de suite et ≥ 300 $ dépensés ; « Augmenter +20 % » si coût par LQ < 100 $ et "
                    "signature ≥ 25 %.")
wt[f"B{WT - 1}"].font = SMALL_I
wt[f"C{WT - 1}"] = None

# monthly
MT = max(w_rows.values()) + 4
m_rows = build_block(MT, "Mois : du 1er jour indiqué en C5 à la fin du mois", M_S, M_E, f"DAY(EOMONTH({M_S},0))", False)
# Google Search reference CPLQ (corpo + mariage), stored in C{MT-1}
gC, gM = m_rows[CAMP_NAMES[0]], m_rows[CAMP_NAMES[2]]
wt[f"B{MT - 1}"] = "Coût par LQ de Google Search (Corporatif + Mariage), pour la règle ChatGPT :"
wt[f"B{MT - 1}"].font = SMALL_I
wt[f"C{MT - 1}"] = (f'=IFERROR(({cDep}{gC}+{cDep}{gM})/({cLQ}{gC}+{cLQ}{gM}),1E+99)')
wt[f"C{MT - 1}"].number_format = MONEY
wt[f"D{MT - 1}"] = '(1E+99 = aucun LQ Google ce mois-ci)'
wt[f"D{MT - 1}"].font = SMALL_I
note_r = max(m_rows.values()) + 1
wt[f"B{note_r}"] = ("Décision mensuelle (COMPTE-RENDU.md, jour 30) : Couper = ChatGPT avec ≥ 900 $ dépensés et "
                    "moins de 2 LQ ou un coût par LQ > 2 × Google Search (les 30 $ vont au Google Corporatif) · "
                    "Réduire 30 % = aucun LQ ou coût par LQ > 300 $ · Augmenter +20 % (max par semaine) = coût par LQ "
                    "< 100 $ et signature ≥ 25 % · sinon Garder. Moins de 300 $ dépensés = données insuffisantes. "
                    "Le total reste à 300 $/jour : toute hausse est financée par la campagne au pire coût par LQ.")
wt[f"B{note_r}"].font = SMALL_I
wt[f"B{note_r}"].alignment = WRAP_L
wt.merge_cells(f"B{note_r}:Q{note_r + 2}")

# segment block (monthly)
ST = note_r + 4
wt[f"B{ST}"] = "Répartition par segment (mois analysé) — cible : corporatif ≥ 50 % de la valeur"
wt[f"B{ST}"].font = Font(bold=True, size=12, color=NAVY)
seg_heads = ["Segment", "Leads (tous)", "LQ", "Contrats", "Revenu", "% de la valeur", "ROAS (non attribué)"]
for i, h in enumerate(seg_heads):
    c = wt.cell(ST + 1, 2 + i, h)
    c.fill, c.font, c.alignment, c.border = F_HEAD, WHITE_B, WRAP_C, BORDER
SEGS = ["Corporatif", "Municipal / scolaire", "Mariage", "Privé"]
s_first = ST + 2
s_tot = s_first + len(SEGS)
for i, sname in enumerate(SEGS + ["Total"]):
    r = s_first + i
    wt[f"B{r}"] = sname
    if sname != "Total":
        seg = f'{rng("segment")},$B{r},'
        wt[f"C{r}"] = f"=COUNTIFS({seg}{crit(M_S, M_E)})"
        wt[f"D{r}"] = f"=COUNTIFS({seg}{crit(M_S, M_E, ',' + rng('qualifie') + ',1')})"
        wt[f"E{r}"] = f"=COUNTIFS({seg}{crit(M_S, M_E, ',' + rng('gagne') + ',1')})"
        wt[f"F{r}"] = f"=SUMIFS({rng('montant_gagne')},{seg}{crit(M_S, M_E, ',' + rng('gagne') + ',1')})"
        wt[f"H{r}"] = "—"
    else:
        for col_ in "CDEF":
            wt[f"{col_}{r}"] = f"=SUM({col_}{s_first}:{col_}{s_tot - 1})"
        wt[f"H{r}"] = f'=IF({cDep}{m_rows["Total"]}=0,"—",F{r}/{cDep}{m_rows["Total"]})'
        wt[f"H{r}"].number_format = "0.00"
    wt[f"G{r}"] = f'=IF($F${s_tot}=0,"—",F{r}/$F${s_tot})'
    wt[f"F{r}"].number_format = MONEY
    wt[f"G{r}"].number_format = PCT
    for col_ in "BCDEFGH":
        wt[f"{col_}{r}"].border = BORDER
        if sname == "Total":
            wt[f"{col_}{r}"].fill, wt[f"{col_}{r}"].font = F_TOTAL, BOLD
wt.conditional_formatting.add(f"G{s_first}", FormulaRule(formula=[f"AND(ISNUMBER(G{s_first}),G{s_first}>=0.5)"],
                                                          fill=GREEN_FILL, font=GREEN_FONT))
wt.conditional_formatting.add(f"G{s_first}", FormulaRule(formula=[f"AND(ISNUMBER(G{s_first}),G{s_first}<0.5)"],
                                                          fill=RED_FILL, font=RED_FONT))

# process block (weekly)
PT = s_tot + 3
wt[f"B{PT}"] = "Procédure de réponse (semaine analysée, LQ = routes A et B) — suivi-leads.md §10"
wt[f"B{PT}"].font = Font(bold=True, size=12, color=NAVY)
for i, h in enumerate(["Indicateur", "Valeur", "Cible"]):
    c = wt.cell(PT + 1, 2 + i, h)
    c.fill, c.font, c.alignment, c.border = F_HEAD, WHITE_B, WRAP_C, BORDER
rec, dl, q, ex = rng("date_heure_reception"), rng("delai_min"), rng("qualifie"), rng("est_exemple")
r1, r2, r3 = PT + 2, PT + 3, PT + 4
wt[f"B{r1}"] = "Délai médian de 1er contact (min)"
wt[f"C{r1}"] = ArrayFormula(
    f"C{r1}",
    f'=IFERROR(MEDIAN(IF(({rec}>={W_S})*({rec}<{W_E})*({q}=1)*ISNUMBER({dl})*({ex}<={EXC}),{dl})),"—")')
wt[f"D{r1}"] = "5 minutes au maximum (heures d'ouverture)"
wt[f"B{r2}"] = "% des LQ contactés en 5 min ou moins"
wt[f"C{r2}"] = (f'=IF({cLQ}{w_rows["Total"]}=0,"—",COUNTIFS({crit(W_S, W_E, "," + q + ",1," + dl + ",\"<=5\"")})'
                f'/{cLQ}{w_rows["Total"]})')
wt[f"D{r2}"] = "Le plus près de 100 %"
wt[f"B{r3}"] = "% des LQ contactés en plus de 15 min (ou pas encore)"
wt[f"C{r3}"] = (f'=IF({cLQ}{w_rows["Total"]}=0,"—",1-COUNTIFS({crit(W_S, W_E, "," + q + ",1," + dl + ",\"<=15\"")})'
                f'/{cLQ}{w_rows["Total"]})')
wt[f"D{r3}"] = "0 % (au-dessus de 15 min : corriger la procédure avant d'augmenter le budget)"
for r in (r1, r2, r3):
    for col_ in "BCD":
        wt[f"{col_}{r}"].border = BORDER
wt[f"C{r1}"].number_format = "0"
wt[f"C{r2}"].number_format = PCT
wt[f"C{r3}"].number_format = PCT
wt.conditional_formatting.add(f"C{r1}", FormulaRule(formula=[f'AND(ISNUMBER(C{r1}),C{r1}>{T["delai"]})'],
                                                    fill=RED_FILL, font=RED_FONT))
wt.conditional_formatting.add(f"C{r1}", FormulaRule(formula=[f'AND(ISNUMBER(C{r1}),C{r1}<={T["delai"]})'],
                                                    fill=GREEN_FILL, font=GREEN_FONT))

wt.column_dimensions["A"].width = 2
wt.column_dimensions["B"].width = 34
for col_ in BCOLS[1:-1]:
    wt.column_dimensions[col_].width = 12
wt.column_dimensions["C"].width = 14
wt.column_dimensions["D"].width = 13
wt.column_dimensions[cDec].width = 28
wt.column_dimensions["R"].width = 12
wt.column_dimensions["S"].width = 10
wt.freeze_panes = "C4"
wt.sheet_properties.tabColor = "548235"

# ---------------------------------------------------------------- sheet: Import conversions Google
wg = wb.create_sheet("Import conversions Google", 3)
G_HEAD = ["Google Click ID", "Conversion Name", "Conversion Time", "Conversion Value", "Conversion Currency"]
CTRL = ["lead_id", "Campagne", "consent_mesure", "date_depot (brute)"]
for i, h in enumerate(G_HEAD, start=1):
    c = wg.cell(1, i, h)
    c.fill, c.font, c.alignment, c.border = F_HEAD, WHITE_B, WRAP_C, BORDER
for i, h in enumerate(CTRL, start=7):
    c = wg.cell(1, i, h)
    c.fill, c.font, c.alignment, c.border = PatternFill("solid", fgColor="595959"), WHITE_B, WRAP_C, BORDER
G_ROWS = 300
for r in range(2, 2 + G_ROWS):
    k = r - 1
    m = f"MATCH({k},{rng('rang_import')},0)"
    wg[f"G{r}"] = f'=IFERROR(INDEX({rng("lead_id")},{m}),"")'
    wg[f"H{r}"] = f'=IF(G{r}="","",INDEX({rng("campagne")},{m}))'
    wg[f"I{r}"] = f'=IF(G{r}="","",INDEX({rng("consent_mesure")},{m})&"")'
    wg[f"J{r}"] = f'=IF(G{r}="","",INDEX({rng("date_depot")},{m}))'
    wg[f"J{r}"].number_format = DT
    d = f"J{r}"
    y = f"YEAR({d})"
    dst = (f"AND({d}>=DATE({y},3,8)+MOD(8-WEEKDAY(DATE({y},3,8)),7)+2/24,"
           f"{d}<DATE({y},11,1)+MOD(8-WEEKDAY(DATE({y},11,1)),7)+2/24)")
    wg[f"A{r}"] = f'=IF(G{r}="","",INDEX({rng("gclid")},{m}))'
    wg[f"B{r}"] = f'=IF(G{r}="","","depot_paye")'
    wg[f"C{r}"] = (f'=IF(G{r}="","",YEAR({d})&"-"&RIGHT("0"&MONTH({d}),2)&"-"&RIGHT("0"&DAY({d}),2)&" "'
                   f'&RIGHT("0"&HOUR({d}),2)&":"&RIGHT("0"&MINUTE({d}),2)&":"&RIGHT("0"&SECOND({d}),2)'
                   f'&IF({dst},"-04:00","-05:00"))')
    wg[f"D{r}"] = f'=IF(G{r}="","",INDEX({rng("montant_gagne")},{m}))'
    wg[f"E{r}"] = f'=IF(G{r}="","","CAD")'
    wg[f"D{r}"].number_format = "0.00"
    for col_ in "ABCDE":
        wg[f"{col_}{r}"].border = BORDER
    for col_ in "GHIJ":
        wg[f"{col_}{r}"].fill = F_CALC
        wg[f"{col_}{r}"].border = BORDER
for col_, w in zip("ABCDEFGHIJ", (26, 16, 26, 16, 18, 3, 14, 26, 12, 17)):
    wg.column_dimensions[col_].width = w
wg.column_dimensions["L"].width = 70
notes = [
    "Comment utiliser cette feuille",
    "1. Ne rien saisir ici : les lignes se remplissent seules à partir de la feuille Leads.",
    "2. Une ligne apparaît quand un lead (non EXEMPLE) est à l'étape « Dépôt payé (gagné) » ou plus, "
    "avec un gclid, une date de dépôt et un montant gagné.",
    "3. Colonnes A à E = format d'import hors ligne de Google Ads (Google Click ID, Conversion Name, "
    "Conversion Time, Conversion Value, Conversion Currency). Heure de l'Est, décalage -05:00 / -04:00 (heure avancée) calculé.",
    "4. Exporter A1:E de cette feuille seulement (Fichier → Enregistrer sous → CSV), ou copier dans la "
    "feuille Google Sheets du téléversement planifié (Objectifs → Conversions → Téléversements → Planifier).",
    "5. Conversion Name doit être identique à l'action Google Ads : depot_paye. Délai maximal : 90 jours après le clic.",
    "6. Courriel et téléphone hachés (SHA-256) : tracking-setup.md §3 les prévoit seulement si consent_mesure = oui. "
    "Excel ne sait pas hacher ; ajouter Email / Phone Number hachés dans Google Sheets (Apps Script) ou dans Make, jamais en clair.",
    "7. Colonnes G à J = contrôle (ne pas téléverser). Microsoft : même logique avec le msclkid (import hebdomadaire).",
]
for i, t in enumerate(notes):
    c = wg.cell(1 + i * 2 if i else 1, 12, t)
    c.alignment = WRAP_L
    c.font = BOLD if i == 0 else Font(size=10)
wg.freeze_panes = "A2"
wg.sheet_properties.tabColor = "2F75B5"

# ---------------------------------------------------------------- sheet: Mode d'emploi
wm = wb.create_sheet("Mode d'emploi", 4)
wm.column_dimensions["A"].width = 3
wm.column_dimensions["B"].width = 120
lines = [
    ("CRM et indicateurs Évenox — mode d'emploi", TITLE),
    ("Sources : operations/tracking-setup.md (pipeline, colonnes, indicateurs), formulaire-qualification.md "
     "(routes, pointage), suivi-leads.md (procédure), COMPTE-RENDU.md (budget 300 $/jour, règles couper / augmenter).", SMALL_I),
    ("", None),
    ("Avant de commencer", BOLD),
    ("• Supprimer les 5 lignes EXEMPLE (surlignées en jaune) de la feuille Leads. Pour les voir dans le tableau de bord, "
     "mettre « Oui » en C6 de « Tableau de bord ». Les lignes EXEMPLE ne vont jamais dans l'import Google.", None),
    ("• Ne pas saisir dans les colonnes grises : ce sont des formules (délai, pointage, campagne, route calculée, valeur, etc.). "
     "Pour supprimer un lead, effacer le contenu des cellules de saisie seulement, ou supprimer la ligne entière.", None),
    ("", None),
    ("Feuille « Leads » (1 ligne par demande)", BOLD),
    ("• Colonnes A à AM = colonnes de la feuille « Pipeline » de tracking-setup.md §9, dans le même ordre ; ligne 2 = nom technique "
     "(mêmes noms que les propriétés HubSpot et les champs du formulaire).", None),
    ("• Dates : AAAA-MM-JJ HH:MM. Le délai de 1er contact est calculé en minutes ; la ligne devient rouge si le délai dépasse "
     "15 min sur une route A ou B.", None),
    ("• type_evenement, type_client, nb_invites et budget_declare utilisent les codes du formulaire (corpo_party, 2500_4999…) : "
     "libellés dans la feuille Listes.", None),
    ("• Route : celle envoyée par le formulaire (A, B, C, D). Si elle est vide, la route calculée selon les règles est utilisée. "
     "Orange = la route saisie diffère des règles du formulaire.", None),
    ("• Pointage (0 à 100) = section 7 du formulaire (budget, type de client, invités, date, rôle, échéance). Plus de 70 : rappel "
     "immédiat par le propriétaire ; 40 à 70 : 15 min par la personne de garde.", None),
    ("• Étape : mettre à jour immédiatement après chaque contact. Perdu ou Disqualifié exige une raison (cellule rouge sinon).", None),
    ("• Gagné = « Dépôt payé (gagné) », « Événement livré » ou « Avis obtenu ». Inscrire alors montant_gagne et date_depot.", None),
    ("• Campagne : déduite de utm_campaign (gads_corpo_*, gads_mariage_*, gads_marque, meta_*, msft_*, chatgpt_*). Pour un appel ou une "
     "référence sans UTM, choisir la source dans « Source / campagne (forcer) ».", None),
    ("• Valeur : montant gagné si gagné, sinon montant proposé × probabilité de l'étape (10/20/40/50/70/100 %). "
     "Valeur de conversion provisoire : A = 400 $, B = 150 $, C et D = 0 $.", None),
    ("", None),
    ("Feuille « Dépenses »", BOLD),
    ("• Chaque matin, saisir la dépense de la veille par campagne (cellules jaunes). Budget cible par campagne en ligne 3 "
     "(100 / 80 / 60 / 30 / 20 / 10 = 300 $/jour) : l'ajuster lors de la bascule saisonnière ou d'une décision Réduire / Augmenter.", None),
    ("• Écart quotidien > +10 % en rouge, < -10 % en orange. Lignes 4 à 6 : cumul, budget des jours saisis et écart.", None),
    ("", None),
    ("Feuille « Tableau de bord » (revue chaque lundi, 9 h)", BOLD),
    ("• C4 = lundi de la semaine analysée ; C5 = 1er jour du mois analysé (par défaut : semaine et mois en cours).", None),
    ("• Vert / rouge selon les cibles : coût par LQ ≤ 150 $, % qualifiés ≥ 40 %, signature ≥ 25 %, panier ≥ 1 500 $, ROAS ≥ 2,5. "
     "Les cibles (colonne I) se modifient après 60 jours.", None),
    ("• Décision : Couper / Réduire 30 % / Garder / Augmenter +20 %. Hebdo = règle de tracking-setup.md §10 ; mensuelle = règle du jour 30 "
     "de COMPTE-RENDU.md. Le total reste à 300 $/jour : une hausse est financée par la campagne au pire coût par LQ.", None),
    ("• Les contrats d'un mois se jugent à 60 à 90 jours (surtout en mariage) : relire les mois précédents avant de couper.", None),
    ("", None),
    ("Feuille « Import conversions Google »", BOLD),
    ("• Remplie automatiquement pour chaque dépôt payé avec gclid. Exporter A:E en CSV ou coller dans la feuille Google Sheets du "
     "téléversement planifié, chaque jour. Ne jamais téléverser les lignes EXEMPLE ni les colonnes de contrôle.", None),
    ("", None),
    ("Confidentialité (Loi 25, LCAP)", BOLD),
    ("• Ce fichier contient des renseignements personnels : le garder dans un espace à accès restreint. Respecter desabonne = oui partout. "
     "Supprimer ou anonymiser les leads perdus sans relation d'affaires après 24 mois.", None),
]
for i, (t, f) in enumerate(lines, start=1):
    c = wm.cell(i, 2, t)
    c.alignment = WRAP_L
    if f:
        c.font = f
wm.sheet_properties.tabColor = "7030A0"

# order: Leads, Dépenses, Tableau de bord, Import, Mode d'emploi, Listes
wb._sheets = [wb["Leads"], wb["Dépenses"], wb["Tableau de bord"], wb["Import conversions Google"],
              wb["Mode d'emploi"], wb["Listes"]]
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("OK", OUT, "leads cols:", len(COLS))
