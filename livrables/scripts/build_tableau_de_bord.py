#!/usr/bin/env python3
"""Génère livrables/TABLEAU-DE-BORD-GOOGLE-ADS.xlsx (tableau de bord hebdo Évenox).

Usage : python3 livrables/scripts/build_tableau_de_bord.py
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "TABLEAU-DE-BORD-GOOGLE-ADS.xlsx"

# ---------------------------------------------------------------- styles
MONEY = '#,##0.00 "$"'
MONEY0 = '#,##0 "$"'
INT = '#,##0'
DEC1 = '#,##0.0'
PCT = '0.0 %'
PCT0 = '0 %'
ROAS = '0.0":1"'
DATE = 'yyyy-mm-dd'
MONTH = '[$-C0C]mmmm yyyy'

YELLOW = PatternFill("solid", fgColor="FFF6C7")
ORANGE = PatternFill("solid", fgColor="F8CBAD")
HEAD = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="D9E1F2")
GREYF = PatternFill("solid", fgColor="EDEDED")
CALC = PatternFill("solid", fgColor="F2F2F2")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="1F3A5F")
NOTE = Font(italic=True, size=9, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CWRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")

GREEN_CF = PatternFill("solid", fgColor="C6EFCE", bgColor="C6EFCE")
YELLOW_CF = PatternFill("solid", fgColor="FFEB9C", bgColor="FFEB9C")
RED_CF = PatternFill("solid", fgColor="FFC7CE", bgColor="FFC7CE")
GREY_CF = PatternFill("solid", fgColor="EDEDED", bgColor="EDEDED")

CAMPAGNES = [
    # nom, budget plafond/jour, CPC max de référence, statut au lancement
    ("S-FR | Marque", 10, 2.50, "À activer"),
    ("S-FR | Corporatif & Fêtes", 170, 4.00, "À activer"),
    ("S-FR | Événements privés", 120, 3.00, "À activer"),
    ("S-EN | Corporate Montréal", 30, 5.00, "En pause (réserve)"),
    ("S-FR | Québec-Lévis (test)", 20, 3.00, "En pause (réserve)"),
    ("Demand Gen", 50, None, "Levier 2 / 4 (section 7)"),
    ("Microsoft Ads", 18, None, "Levier 3 (section 7)"),
]
NC = len(CAMPAGNES)
DG_IDX, MS_IDX = 5, 6
SEARCH_IDX = range(0, 5)

FIRST_WEEK = date(2026, 10, 12)
LAST_WEEK = date(2027, 9, 27)
MONTHS = [date(2026 + (9 + i) // 12, (9 + i) % 12 + 1, 1) for i in range(12)]  # oct. 2026 → sept. 2027

SH = "'Saisie hebdo'"
MAXR = 2000  # plage des formules (permet d'ajouter des lignes)


def hdr(ws, row, labels, fill=HEAD, font=WHITE_B, height=45):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row, c, lab)
        cell.fill, cell.font, cell.alignment, cell.border = fill, font, CWRAP, BOX
    ws.row_dimensions[row].height = height


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[L(i)].width = w


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def div(num, den):
    return f'=IF(N({den})=0,"",{num}/{den})'


wb = Workbook()

# =================================================================== CIBLES
cb = wb.active
cb.title = "Cibles"
ws_saisie = wb.create_sheet("Saisie hebdo")
ws_synth = wb.create_sheet("Synthèse mensuelle")
ws_reg = wb.create_sheet("Règles")
ws_mode = wb.create_sheet("Mode d'emploi")
wb.move_sheet("Cibles", offset=3)  # ordre : Saisie, Synthèse, Règles, Cibles, Mode d'emploi

cb["A1"] = "Cibles et hypothèses — modifiables"
cb["A1"].font = TITLE
cb["A2"] = ("Cellules jaunes = à ajuster. La feuille « Règles » lit ces valeurs : "
            "changez-les ici, jamais dans les formules.")
cb["A2"].font = NOTE

hdr(cb, 4, ["Panier moyen par type d'événement", "Panier moyen ($)", "Source / note"], height=30)
paniers = [
    ("Party de Bureau", 1995, "Forfait tout inclus — evenox.ca, 7 oct. 2026"),
    ("5 à 7 d'équipe", 1195, "Forfait tout inclus — evenox.ca"),
    ("Gala Signature", 2495, "Forfait tout inclus — evenox.ca"),
    ("Décor WOW", 899, "Mariage et événement — evenox.ca"),
    ("Soirée Signature", 1449, "Mariage et événement — evenox.ca"),
    ("Photobooth Signature 3 h", 799, "Location photobooth — evenox.ca"),
]
r = 5
for t, v, s in paniers:
    cb.cell(r, 1, t)
    c = cb.cell(r, 2, v)
    c.number_format, c.fill = MONEY, YELLOW
    cb.cell(r, 3, s).font = NOTE
    r += 1
cb.cell(11, 1, "Panier moyen de référence").font = BOLD
c = cb.cell(11, 2, "=AVERAGE(B5:B10)")
c.number_format, c.fill, c.font = MONEY, YELLOW, BOLD
cb.cell(11, 3, "Moyenne simple des 6 forfaits. Remplacez par votre panier réel "
               "(Valeur des contrats ÷ Contrats) après 10 contrats.").font = NOTE
name(wb, "Panier_ref", "Cibles!$B$11")

hdr(cb, 13, ["Rentabilité", "Valeur", "Note"], height=30)
rent = [
    ("Marge brute", 0.50, PCT0, True, "À confirmer avec ta comptabilité.", "Marge_brute"),
    ("Plafond coût par lead", 90, MONEY, True, "Au-delà : baisser les CPC ou resserrer les mots-clés.", "Plafond_CPL"),
    ("Plafond coût par contrat (% du panier)", 0.15, PCT0, True, "15 % du panier moyen de référence.", "Plafond_contrat_pct"),
    ("Plafond coût par contrat ($)", "=B16*Panier_ref", MONEY, False, "Calculé : % × panier de référence.", "Plafond_contrat"),
    ("ROAS minimum (cible)", 3, ROAS, True, "Valeur des contrats ÷ coût publicitaire.", "ROAS_min"),
    ("ROAS point mort", "=1/Marge_brute", ROAS, False, "Calculé : 1 ÷ marge brute (2:1 à 50 % de marge).", "ROAS_point_mort"),
]
r = 14
for lab, v, fmt, editable, note, nm in rent:
    cb.cell(r, 1, lab)
    c = cb.cell(r, 2, v)
    c.number_format = fmt
    c.fill = YELLOW if editable else CALC
    cb.cell(r, 3, note).font = NOTE
    name(wb, nm, f"Cibles!$B${r}")
    r += 1

hdr(cb, 21, ["Seuils des règles (plan, sections 5 à 7)", "Valeur", "Section du plan"], height=30)
seuils = [
    ("Conversions sur 30 jours → Max. conversions", 15, INT, "5 — Apprentissage", "S_MaxConv"),
    ("Conversions sur 30 jours → CPA cible", 30, INT, "5 — Efficacité", "S_CPAcible"),
    ("Majoration du CPA cible (CPA 30 j + …)", 0.10, PCT0, "5 — Efficacité", "S_CPAmaj"),
    ("Leads qualifiés / 30 jours → « Lead qualifié » en conversion principale", 15, INT, "5 — Qualité", "S_Qualite"),
    ("Contrats (conversions avec valeur) / 30 jours → Max. valeur", 30, INT, "5 — Valeur", "S_Valeur"),
    ("Coupe-circuit : coût par lead qualifié sur 14 jours", 150, MONEY, "5 — Coupe-circuit", "S_CoupeCPLQ"),
    ("Coupe-circuit : CPC moyen > N × CPC de référence (7 jours)", 2, DEC1, "5 — Coupe-circuit", "S_CoupeCPC"),
    ("% de leads qualifiés minimum", 0.50, PCT0, "6 — Routine", "S_PctQual"),
    ("Part d'impressions perdue (classement) maximum", 0.30, PCT0, "7 — Levier 1", "S_PerdueClass"),
    ("Hausse des CPC max. par semaine", 0.20, PCT0, "7 — Levier 1", "S_HausseCPC"),
    ("Demand Gen prospection : dépense max. sans lead qualifié", 1500, MONEY, "7 — Levier 2 (remarketing : 1 000 $)", "S_DG"),
    ("Microsoft Ads : coût par lead max. vs Google (×)", 1.5, DEC1, "7 — Levier 3", "S_MS"),
]
r = 22
for lab, v, fmt, sec, nm in seuils:
    cb.cell(r, 1, lab)
    c = cb.cell(r, 2, v)
    c.number_format, c.fill = fmt, YELLOW
    cb.cell(r, 3, sec).font = NOTE
    name(wb, nm, f"Cibles!$B${r}")
    r += 1

CAMP_ROW0 = 36
hdr(cb, 35, ["Campagne", "Budget plafond / jour", "CPC max. de référence", "Statut au lancement"], height=30)
for i, (n, b, cpc, st) in enumerate(CAMPAGNES):
    rr = CAMP_ROW0 + i
    cb.cell(rr, 1, n)
    c = cb.cell(rr, 2, b); c.number_format, c.fill = MONEY0, YELLOW
    c = cb.cell(rr, 3, cpc); c.number_format, c.fill = MONEY, YELLOW
    cb.cell(rr, 4, st)
name(wb, "Liste_campagnes", f"Cibles!$A${CAMP_ROW0}:$A${CAMP_ROW0 + NC - 1}")
cb.cell(CAMP_ROW0 + NC, 1, "Les noms doivent être identiques à ceux de Google Ads pour faciliter la saisie. "
                           "CPC de référence vide = règle CPC non applicable.").font = NOTE
for row in cb.iter_rows(min_row=4, max_row=CAMP_ROW0 + NC - 1, max_col=4):
    for c in row:
        if c.row not in (12, 20, 34):
            c.border = BOX
widths(cb, [62, 20, 22, 46])
for rr in range(4, CAMP_ROW0 + NC):
    cb.cell(rr, 3).alignment = Alignment(wrap_text=True, vertical="center")
cb.freeze_panes = "A4"

# =================================================================== SAISIE HEBDO
ws = ws_saisie
INPUTS = [
    ("Coût", MONEY), ("Impressions", INT), ("Clics", INT),
    ("Conversions Google (formulaire + appels)", DEC1),
    ("Part d'impr. perdue (budget) %", PCT), ("Part d'impr. perdue (classement) %", PCT),
    ("Leads reçus (Notion)", INT), ("Leads qualifiés A/B/C", INT),
    ("Soumissions envoyées", INT), ("Contrats gagnés", INT), ("Valeur des contrats $", MONEY),
]
# colonnes : A semaine, B mois, C campagne, D remarque, E..O saisies, P..Y calculs
cE, cF, cG, cH, cI, cJ, cK, cLq, cM, cN, cO = [L(5 + i) for i in range(11)]
CALCS = [
    ("CTR", PCT, lambda r: div(f"G{r}", f"F{r}")),
    ("CPC moyen", MONEY, lambda r: div(f"E{r}", f"G{r}")),
    ("Taux de conversion", PCT, lambda r: div(f"H{r}", f"G{r}")),
    ("CPA Google (coût ÷ conv.)", MONEY, lambda r: div(f"E{r}", f"H{r}")),
    ("Coût par lead", MONEY, lambda r: div(f"E{r}", f"K{r}")),
    ("Coût par lead qualifié", MONEY, lambda r: div(f"E{r}", f"L{r}")),
    ("Coût par contrat", MONEY, lambda r: div(f"E{r}", f"N{r}")),
    ("% qualifiés", PCT0, lambda r: div(f"L{r}", f"K{r}")),
    ("Taux de conclusion (contrats ÷ soumissions)", PCT0, lambda r: div(f"N{r}", f"M{r}")),
    ("ROAS (valeur ÷ coût)", ROAS, lambda r: div(f"O{r}", f"E{r}")),
]

ws["A1"] = "Saisie hebdomadaire — Google Ads × Notion"
ws["A1"].font = TITLE
ws["A2"] = ("Chaque lundi : remplir les cellules jaunes de la semaine précédente (lundi au dimanche). "
            "Colonnes grises = calcul automatique, ne pas modifier. Le mois est celui du lundi de la semaine.")
ws["A2"].font = NOTE
for c in range(1, 5):
    ws.cell(3, c).fill = SUB
ws.cell(3, 1, "REPÈRES").font = BOLD
for c in range(5, 16):
    ws.cell(3, c).fill = YELLOW
ws.cell(3, 5, "SAISIE (cellules jaunes) — Google Ads : Coût → Part perdue · Notion : Leads → Valeur").font = BOLD
for c in range(16, 26):
    ws.cell(3, c).fill = CALC
ws.cell(3, 16, "CALCULS AUTOMATIQUES").font = BOLD

hdr(ws, 4, ["Semaine (lundi)", "Mois", "Campagne", "Remarque"] + [n for n, _ in INPUTS] + [n for n, _, _ in CALCS],
    height=60)

EXEMPLES = {
    ("S-FR | Marque"): [45.30, 210, 38, 6, 0, 0.12, 5, 4, 3, 1, 1995],
    ("S-FR | Corporatif & Fêtes"): [612.40, 3150, 168, 9, 0.05, 0.38, 8, 5, 4, 1, 2495],
}
r = 5
wk = FIRST_WEEK
while wk <= LAST_WEEK:
    for n, *_ in CAMPAGNES:
        c = ws.cell(r, 1, wk); c.number_format = DATE
        c = ws.cell(r, 2, f"=DATE(YEAR(A{r}),MONTH(A{r}),1)"); c.number_format = MONTH
        ws.cell(r, 3, n)
        for i, (_, fmt) in enumerate(INPUTS):
            c = ws.cell(r, 5 + i)
            c.fill, c.number_format = YELLOW, fmt
        if wk == FIRST_WEEK and n in EXEMPLES:
            ws.cell(r, 4, "EXEMPLE — à effacer").fill = ORANGE
            ws.cell(r, 4).font = Font(bold=True, color="C00000")
            for i, v in enumerate(EXEMPLES[n]):
                ws.cell(r, 5 + i, v)
        for j, (_, fmt, f) in enumerate(CALCS):
            c = ws.cell(r, 16 + j, f(r))
            c.number_format, c.fill = fmt, CALC
        if n == CAMPAGNES[-1][0]:
            for cc in range(1, 26):
                ws.cell(r, cc).border = Border(bottom=Side(style="thin", color="808080"))
        r += 1
    wk += timedelta(days=7)
LAST_DATA = r - 1

widths(ws, [12, 14, 27, 20] + [12, 12, 9, 14, 12, 13, 11, 11, 12, 10, 13] + [9, 10, 11, 12, 11, 12, 12, 10, 13, 10])
ws.freeze_panes = "E5"
ws.auto_filter.ref = f"A4:{L(25)}{LAST_DATA}"

dv_camp = DataValidation(type="list", formula1="Liste_campagnes", allow_blank=True,
                         errorTitle="Campagne inconnue", error="Choisir une campagne de la liste (feuille Cibles).")
dv_pos = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        errorTitle="Valeur invalide", error="Entrer un nombre positif ou zéro.")
dv_int = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        errorTitle="Valeur invalide", error="Entrer un nombre entier positif ou zéro.")
dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        errorTitle="Pourcentage invalide", error="Entrer un pourcentage entre 0 % et 100 % (ex. : 12 %).")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True,
                         errorTitle="Date invalide", error="Entrer la date du lundi (aaaa-mm-jj).")
for dv in (dv_camp, dv_pos, dv_int, dv_pct, dv_date):
    ws.add_data_validation(dv)
dv_date.add(f"A5:A{MAXR}")
dv_camp.add(f"C5:C{MAXR}")
dv_pos.add(f"E5:E{MAXR}"); dv_pos.add(f"H5:H{MAXR}"); dv_pos.add(f"O5:O{MAXR}")
dv_int.add(f"F5:G{MAXR}"); dv_int.add(f"K5:N{MAXR}")
dv_pct.add(f"I5:J{MAXR}")
# contrôle : qualifiés ≤ leads
ws.conditional_formatting.add(f"L5:L{MAXR}", FormulaRule(formula=["AND(L5<>\"\",L5>N(K5))"], fill=RED_CF))
ws.conditional_formatting.add(f"N5:N{MAXR}", FormulaRule(formula=["AND(N5<>\"\",N5>N(M5))"], fill=RED_CF))


def rng(col):
    return f"{SH}!${col}$5:${col}${MAXR}"


# =================================================================== SYNTHÈSE MENSUELLE
sy = ws_synth
sy["A1"] = "Synthèse mensuelle par campagne"
sy["A1"].font = TITLE
sy["A2"] = ("Calculée automatiquement à partir de « Saisie hebdo » (une semaine compte dans le mois de son lundi). "
            "Parts d'impressions perdues = moyenne simple des semaines saisies. Ne rien saisir ici.")
sy["A2"].font = NOTE
SY_IN = ["Coût", "Impressions", "Clics", "Conversions Google", "Part perdue budget (moy.)",
         "Part perdue classement (moy.)", "Leads reçus", "Leads qualifiés A/B/C", "Soumissions",
         "Contrats gagnés", "Valeur des contrats"]
hdr(sy, 4, ["Mois", "Campagne"] + SY_IN + [n for n, _, _ in CALCS], height=60)
SRC = ["E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O"]
FMTS = [f for _, f in INPUTS]
# colonnes synthèse : C..M = saisies agrégées ; N..W = KPI
SYC = {s: L(3 + i) for i, s in enumerate(SRC)}  # E->C, F->D, ...


def sy_kpis(r):
    m = {k: f"{v}{r}" for k, v in SYC.items()}
    return [
        div(m["G"], m["F"]), div(m["E"], m["G"]), div(m["H"], m["G"]), div(m["E"], m["H"]),
        div(m["E"], m["K"]), div(m["E"], m["L"]), div(m["E"], m["N"]), div(m["L"], m["K"]),
        div(m["N"], m["M"]), div(m["O"], m["E"]),
    ]


def sy_row(r, crit, label_month, label_camp, total=False):
    if isinstance(label_month, date):
        c = sy.cell(r, 1, label_month); c.number_format = MONTH
    else:
        sy.cell(r, 1, label_month)
    sy.cell(r, 2, label_camp)
    for i, s in enumerate(SRC):
        if s in ("I", "J"):
            f = f'=IFERROR(AVERAGEIFS({rng(s)}{crit}),"")' if crit else f'=IFERROR(AVERAGE({rng(s)}),"")'
        else:
            f = f"=SUMIFS({rng(s)}{crit})" if crit else f"=SUM({rng(s)})"
        c = sy.cell(r, 3 + i, f)
        c.number_format = FMTS[i]
    for j, f in enumerate(sy_kpis(r)):
        c = sy.cell(r, 14 + j, f)
        c.number_format = CALCS[j][1]
    if total:
        for cc in range(1, 24):
            sy.cell(r, cc).fill = SUB
            sy.cell(r, cc).font = BOLD
            sy.cell(r, cc).border = Border(bottom=Side(style="medium", color="1F3A5F"))


r = 5
for m in MONTHS:
    for n, *_ in CAMPAGNES:
        sy_row(r, f",{rng('B')},$A{r},{rng('C')},$B{r}", m, n)
        r += 1
    sy_row(r, f",{rng('B')},$A{r}", m, "Total du compte", total=True)
    r += 1
r += 1
sy.cell(r, 1, "Cumul octobre 2026 à septembre 2027").font = TITLE
r += 1
for n, *_ in CAMPAGNES:
    sy_row(r, f",{rng('C')},$B{r}", "Cumul", n)
    r += 1
sy_row(r, "", "Cumul", "Total du compte", total=True)
SY_LAST = r
widths(sy, [15, 27] + [12, 12, 9, 12, 12, 13, 10, 12, 12, 10, 13] + [9, 10, 11, 12, 11, 12, 12, 10, 13, 10])
sy.freeze_panes = "C5"
sy.auto_filter.ref = f"A4:W{SY_LAST - NC - 3}"
# repères visuels : coût par lead / ROAS
sy.conditional_formatting.add(f"R5:R{SY_LAST}", FormulaRule(formula=['AND(ISNUMBER(R5),R5>Plafond_CPL)'], fill=RED_CF))
sy.conditional_formatting.add(f"S5:S{SY_LAST}", FormulaRule(formula=['AND(ISNUMBER(S5),S5>S_CoupeCPLQ)'], fill=RED_CF))
sy.conditional_formatting.add(f"W5:W{SY_LAST}", FormulaRule(formula=['AND(ISNUMBER(W5),W5<ROAS_point_mort)'], fill=RED_CF))
sy.conditional_formatting.add(f"W5:W{SY_LAST}", FormulaRule(formula=['AND(ISNUMBER(W5),W5<ROAS_min)'], fill=YELLOW_CF))
sy.conditional_formatting.add(f"W5:W{SY_LAST}", FormulaRule(formula=['AND(ISNUMBER(W5),W5>=ROAS_min)'], fill=GREEN_CF))

# =================================================================== RÈGLES
rg = ws_reg
rg["A1"] = "Règles de décision — statut automatique"
rg["A1"].font = TITLE
rg["A2"] = ("Basé sur les 4 dernières semaines saisies (≈ 30 jours) et, pour les contrats et le ROAS, sur le cumul "
            "(les contrats arrivent avec un délai). Seuils modifiables dans « Cibles ». "
            "Rouge = ALERTE (agir aujourd'hui) · Jaune = À FAIRE · Vert = OK · Gris = pas de données / sans objet.")
rg["A2"].font = NOTE
rg["A2"].alignment = WRAP
rg.row_dimensions[2].height = 30

params = [
    ("Dernière semaine saisie (auto)", f"=SUMPRODUCT(MAX(({rng('E')}<>\"\")*{rng('A')}))", None),
    ("Forcer une autre semaine (optionnel, un lundi)", None, None),
    ("Semaine de référence utilisée", '=IF(B5<>"",B5,B4)', None),
    ("Période 4 semaines (≈ 30 jours)", "=B6-21", "=B6+6"),
    ("Période 2 semaines (14 jours)", "=B6-7", "=B6+6"),
]
for i, (lab, f1, f2) in enumerate(params):
    rr = 4 + i
    rg.cell(rr, 1, lab).font = BOLD
    c = rg.cell(rr, 2, f1); c.number_format = DATE; c.border = BOX
    if f2:
        rg.cell(rr, 3, f2).number_format = DATE
        rg.cell(rr, 3).border = BOX
rg["B5"].fill = YELLOW
rg["C4"] = '=IF(B4=0,"Aucune donnée saisie pour l\'instant","")'
rg["C4"].font = Font(bold=True, color="C00000")
rg.add_data_validation(DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)",
                                      allow_blank=True, sqref="B5"))
name(wb, "Sem_ref", "'Règles'!$B$6")
name(wb, "Debut4", "'Règles'!$B$7")
name(wb, "Debut2", "'Règles'!$B$8")

# --- tableau d'indicateurs
IND_H = 11
IND0 = 12
IND_COLS = [
    # (titre, format, type, source/args)
    ("Coût 4 sem.", MONEY, "sum4", "E"),                    # B
    ("Conversions Google 4 sem.", DEC1, "sum4", "H"),       # C
    ("Leads 4 sem.", INT, "sum4", "K"),                     # D
    ("Qualifiés 4 sem.", INT, "sum4", "L"),                 # E
    ("Contrats 4 sem.", INT, "sum4", "N"),                  # F
    ("Valeur 4 sem.", MONEY, "sum4", "O"),                  # G
    ("Coût par lead 4 sem.", MONEY, "div", ("B", "D")),     # H
    ("% qualifiés 4 sem.", PCT0, "div", ("E", "D")),        # I
    ("Part perdue classement (moy. 4 sem.)", PCT, "avg4", "J"),  # J
    ("CPA Google 4 sem.", MONEY, "div", ("B", "C")),        # K
    ("CPA cible suggéré (+10 %)", MONEY, "cpa", None),      # L
    ("Coût 2 sem.", MONEY, "sum2", "E"),                    # M
    ("Qualifiés 2 sem.", INT, "sum2", "L"),                 # N
    ("Coût par lead qualifié 2 sem.", MONEY, "div", ("M", "N")),  # O
    ("CPC moyen semaine de réf.", MONEY, "cpcw", None),     # P
    ("Coût cumulé", MONEY, "cum", "E"),                     # Q
    ("Leads cumulés", INT, "cum", "K"),                     # R
    ("Qualifiés cumulés", INT, "cum", "L"),                 # S
    ("Contrats cumulés", INT, "cum", "N"),                  # T
    ("Valeur cumulée", MONEY, "cum", "O"),                  # U
    ("Coût par lead cumulé", MONEY, "div", ("Q", "R")),     # V
    ("Coût par contrat cumulé", MONEY, "div", ("Q", "T")),  # W
    ("ROAS cumulé", ROAS, "div", ("U", "Q")),               # X
]
rg.cell(IND_H - 1, 1, "1. Indicateurs par campagne (calcul automatique)").font = TITLE
hdr(rg, IND_H, ["Campagne"] + [t for t, *_ in IND_COLS], height=60)
TOT_S = IND0 + NC        # total Google Search
TOT_A = IND0 + NC + 1    # total compte
for i in range(NC + 2):
    rr = IND0 + i
    if i < NC:
        rg.cell(rr, 1, CAMPAGNES[i][0])
        crit = f",{rng('C')},$A{rr}"
    else:
        rg.cell(rr, 1, "Total Google Search (5 campagnes)" if rr == TOT_S else "Total du compte")
        crit = None
    for j, (t, fmt, kind, arg) in enumerate(IND_COLS):
        col = L(2 + j)
        if crit is None and kind in ("sum4", "sum2", "cum", "avg4", "cpcw"):
            rows = f"{col}{IND0}:{col}{IND0 + 4}" if rr == TOT_S else f"{col}{IND0}:{col}{IND0 + NC - 1}"
            if kind == "avg4":
                f = f'=IFERROR(AVERAGE({rows}),"")'
            elif kind == "cpcw":
                f = (f'=IFERROR(SUMIFS({rng("E")},{rng("A")},Sem_ref)/SUMIFS({rng("G")},{rng("A")},Sem_ref),"")'
                     if rr == TOT_A else "")
            else:
                f = f"=SUM({rows})"
        elif kind == "sum4":
            f = f'=SUMIFS({rng(arg)}{crit},{rng("A")},">="&Debut4,{rng("A")},"<="&Sem_ref)'
        elif kind == "sum2":
            f = f'=SUMIFS({rng(arg)}{crit},{rng("A")},">="&Debut2,{rng("A")},"<="&Sem_ref)'
        elif kind == "avg4":
            f = f'=IFERROR(AVERAGEIFS({rng(arg)}{crit},{rng("A")},">="&Debut4,{rng("A")},"<="&Sem_ref),"")'
        elif kind == "cum":
            f = f'=SUMIFS({rng(arg)}{crit},{rng("A")},"<="&Sem_ref)'
        elif kind == "cpcw":
            f = (f'=IFERROR(SUMIFS({rng("E")}{crit},{rng("A")},Sem_ref)/SUMIFS({rng("G")}{crit},{rng("A")},Sem_ref),"")')
        elif kind == "cpa":
            f = f'=IF(K{rr}="","",K{rr}*(1+S_CPAmaj))'
        elif kind == "div":
            a, b = arg
            f = div(f"{a}{rr}", f"{b}{rr}")
        c = rg.cell(rr, 2 + j, f if f else None)
        c.number_format, c.border = fmt, BOX
    rg.cell(rr, 1).border = BOX
    if i >= NC:
        for cc in range(1, 2 + len(IND_COLS)):
            rg.cell(rr, cc).fill = SUB
            rg.cell(rr, cc).font = BOLD

# --- matrice des règles
RH = TOT_A + 3
rg.cell(RH - 1, 1, "2. Règles du plan — que faire cette semaine ?").font = TITLE
hdr(rg, RH, ["Règle", "Déclencheur (seuil modifiable dans Cibles)", "Action prévue au plan", "Section"]
    + [n for n, *_ in CAMPAGNES], height=45)

NODATA = '"— Pas de données"'
SO = '"s.o."'


def ind(col, i):
    return f"${col}${IND0 + i}"


def cpc_ref(i):
    return f"Cibles!$C${CAMP_ROW0 + i}"


def rule_formulas():
    G = lambda col: f"${col}${TOT_S}"  # noqa: E731 — total Google Search
    rules = []

    def add(title, trig, action, section, fn, applies=lambda i: True):
        rules.append((title, trig, action, section, fn, applies))

    not_dg = lambda i: i != DG_IDX  # noqa: E731

    add("Passer en Max. conversions",
        '="Conversions Google sur 4 sem. ≥ "&S_MaxConv',
        "Stratégie Max. conversions, sans cible (si pas déjà fait)", "5",
        lambda i: (f'=IF({ind("B", i)}=0,{NODATA},IF({ind("C", i)}>=S_MaxConv,'
                   f'"À FAIRE : Passer en Max. conversions","OK — rester en Max. clics ("&ROUND({ind("C", i)},0)&"/"&S_MaxConv&")"))'),
        not_dg)
    add("Passer en CPA cible",
        '="Conversions Google sur 4 sem. ≥ "&S_CPAcible',
        "CPA cible = CPA 30 j + 10 %, puis −10 % toutes les 2-3 semaines", "5",
        lambda i: (f'=IF({ind("B", i)}=0,{NODATA},IF({ind("C", i)}>=S_CPAcible,'
                   f'"À FAIRE : Passer en CPA cible ≈ "&ROUND({ind("L", i)},0)&" $","OK — pas encore ("&ROUND({ind("C", i)},0)&"/"&S_CPAcible&")"))'),
        not_dg)
    add("« Lead qualifié » en conversion principale",
        '="Leads qualifiés sur 4 sem. ≥ "&S_Qualite',
        "Lead qualifié (import GCLID) = principale ; formulaire = secondaire", "5",
        lambda i: (f'=IF({ind("B", i)}=0,{NODATA},IF({ind("E", i)}>=S_Qualite,'
                   f'"À FAIRE : Lead qualifié en conversion principale","OK — pas encore ("&{ind("E", i)}&"/"&S_Qualite&")"))'),
        not_dg)
    add("Passer en Max. valeur de conversion",
        '="Contrats (conversions avec valeur) sur 4 sem. ≥ "&S_Valeur',
        "Max. valeur de conversion, puis ROAS cible", "5",
        lambda i: (f'=IF({ind("B", i)}=0,{NODATA},IF({ind("F", i)}>=S_Valeur,'
                   f'"À FAIRE : Passer en Max. valeur de conversion","OK — pas encore ("&{ind("F", i)}&"/"&S_Valeur&")"))'),
        not_dg)
    add("Coupe-circuit — coût par lead qualifié",
        '="Coût par lead qualifié sur 2 sem. > "&S_CoupeCPLQ&" $"',
        "Revenir à l'étape d'enchères précédente", "5",
        lambda i: (f'=IF({ind("M", i)}=0,{NODATA},IF(IF({ind("N", i)}=0,{ind("M", i)}>S_CoupeCPLQ,{ind("O", i)}>S_CoupeCPLQ),'
                   f'"ALERTE : Coupe-circuit — revenir à l\'étape précédente",'
                   f'"OK ("&IF({ind("N", i)}=0,"0 qualifié",ROUND({ind("O", i)},0)&" $ / qualifié")&")"))'))
    add("Coupe-circuit — CPC",
        '="CPC moyen de la semaine > "&S_CoupeCPC&" × CPC de référence"',
        "Revenir à l'étape d'enchères précédente (si 7 jours de suite)", "5",
        lambda i: (f'=IF({cpc_ref(i)}="",{SO},IF({ind("P", i)}="",{NODATA},IF({ind("P", i)}>S_CoupeCPC*{cpc_ref(i)},'
                   f'"ALERTE : Coupe-circuit — CPC "&ROUND({ind("P", i)},2)&" $","OK ("&ROUND({ind("P", i)},2)&" $)")))'),
        lambda i: CAMPAGNES[i][2] is not None)
    add("Qualité des leads",
        '="% qualifiés (A/B/C) sur 4 sem. < "&ROUND(S_PctQual*100,0)&" %"',
        "Revoir mots-clés et négatifs (rapport Termes de recherche)", "6",
        lambda i: (f'=IF({ind("D", i)}=0,{NODATA},IF({ind("I", i)}<S_PctQual,'
                   f'"À FAIRE : Revoir mots-clés/négatifs ("&ROUND({ind("I", i)}*100,0)&" %)","OK ("&ROUND({ind("I", i)}*100,0)&" % qualifiés)"))'))
    add("Coût par lead",
        '="Coût par lead sur 4 sem. > "&Plafond_CPL&" $"',
        "Baisser les CPC max. ou resserrer les mots-clés", "Cibles",
        lambda i: (f'=IF({ind("B", i)}=0,{NODATA},IF({ind("D", i)}=0,IF({ind("B", i)}>Plafond_CPL,'
                   f'"À FAIRE : 0 lead pour "&ROUND({ind("B", i)},0)&" $ — vérifier le suivi","OK — trop tôt"),'
                   f'IF({ind("H", i)}>Plafond_CPL,"À FAIRE : Coût par lead "&ROUND({ind("H", i)},0)&" $ > plafond",'
                   f'"OK ("&ROUND({ind("H", i)},0)&" $ / lead)")))'))
    add("Hausser les CPC (levier 1)",
        '="Part perdue (classement) > "&ROUND(S_PerdueClass*100,0)&" % et coût par lead ≤ "&Plafond_CPL&" $"',
        "Hausser CPC max. +20 %/sem. sur les mots-clés rentables (arrêt si CPA +25 %)", "7",
        lambda i: (f'=IF({ind("B", i)}=0,{NODATA},IF({ind("J", i)}="","— Part perdue non saisie",'
                   f'IF(AND({ind("J", i)}>S_PerdueClass,{ind("D", i)}>0,N({ind("H", i)})<=Plafond_CPL),'
                   f'"À FAIRE : Hausser CPC +"&ROUND(S_HausseCPC*100,0)&" %",'
                   f'IF({ind("J", i)}>S_PerdueClass,"OK — part perdue élevée mais CPA hors cible : ne pas hausser","OK"))))'),
        not_dg)
    add("Coût par contrat",
        '="Coût par contrat cumulé > "&ROUND(Plafond_contrat_pct*100,0)&" % du panier ("&ROUND(Plafond_contrat,0)&" $)"',
        "Revoir ciblage, annonces et suivi des leads (vitesse de rappel)", "Cibles",
        lambda i: (f'=IF({ind("Q", i)}=0,{NODATA},IF({ind("T", i)}=0,IF({ind("Q", i)}>Plafond_contrat,'
                   f'"À FAIRE : 0 contrat pour "&ROUND({ind("Q", i)},0)&" $ — relancer les soumissions","OK — trop tôt"),'
                   f'IF({ind("W", i)}>Plafond_contrat,"ALERTE : "&ROUND({ind("W", i)},0)&" $ par contrat > plafond",'
                   f'"OK ("&ROUND({ind("W", i)},0)&" $ / contrat)")))'))
    add("Rentabilité (ROAS)",
        '="ROAS cumulé < "&ROAS_min&":1 (point mort "&ROUND(ROAS_point_mort,1)&":1)"',
        "Sous le point mort : réduire le budget de la campagne", "Cibles",
        lambda i: (f'=IF({ind("Q", i)}=0,{NODATA},IF(AND({ind("T", i)}=0,{ind("Q", i)}<2*Plafond_contrat),"OK — trop tôt (aucun contrat)",'
                   f'IF({ind("X", i)}<ROAS_point_mort,"ALERTE : ROAS "&ROUND({ind("X", i)},1)&":1 sous le point mort",'
                   f'IF({ind("X", i)}<ROAS_min,"À FAIRE : ROAS "&ROUND({ind("X", i)},1)&":1 sous la cible",'
                   f'"OK (ROAS "&ROUND({ind("X", i)},1)&":1)"))))'))
    add("Pause Demand Gen (leviers 2 et 4)",
        '="Coût cumulé ≥ "&S_DG&" $ et 0 lead qualifié"',
        "Mettre Demand Gen en pause (remarketing : seuil 1 000 $)", "7",
        lambda i: (f'=IF({ind("Q", i)}=0,{NODATA},IF(AND({ind("Q", i)}>=S_DG,{ind("S", i)}=0),"ALERTE : Pause Demand Gen",'
                   f'"OK ("&ROUND({ind("Q", i)},0)&" $ / "&{ind("S", i)}&" qualifié(s))"))'),
        lambda i: i == DG_IDX)
    add("Pause Microsoft Ads (levier 3)",
        '="Coût par lead cumulé > "&S_MS&" × Google Search"',
        "Mettre Microsoft Ads en pause", "7",
        lambda i: (f'=IF({ind("Q", i)}=0,{NODATA},IF({ind("R", i)}=0,"À FAIRE : 0 lead pour "&ROUND({ind("Q", i)},0)&" $",'
                   f'IF({G("V")}="","— Pas de référence Google",IF({ind("V", i)}>S_MS*{G("V")},'
                   f'"ALERTE : Pause Microsoft Ads","OK ("&ROUND({ind("V", i)},0)&" $ vs "&ROUND({G("V")},0)&" $)"))))'),
        lambda i: i == MS_IDX)
    return rules


RULES = rule_formulas()
R0 = RH + 1
for k, (title, trig, action, section, fn, applies) in enumerate(RULES):
    rr = R0 + k
    rg.cell(rr, 1, title).font = BOLD
    rg.cell(rr, 2, trig)
    rg.cell(rr, 3, action)
    rg.cell(rr, 4, section).alignment = CWRAP
    for i in range(NC):
        rg.cell(rr, 5 + i, fn(i) if applies(i) else "s.o.")
    for cc in range(1, 5 + NC):
        c = rg.cell(rr, cc)
        c.border = BOX
        if cc != 4:
            c.alignment = Alignment(wrap_text=True, vertical="center")
    rg.row_dimensions[rr].height = 42
R_LAST = R0 + len(RULES) - 1
area = f"E{R0}:{L(4 + NC)}{R_LAST}"
rg.conditional_formatting.add(area, FormulaRule(formula=[f'LEFT(E{R0},6)="ALERTE"'], fill=RED_CF,
                                                font=Font(bold=True, color="9C0006"), stopIfTrue=True))
rg.conditional_formatting.add(area, FormulaRule(formula=[f'LEFT(E{R0},7)="À FAIRE"'], fill=YELLOW_CF,
                                                font=Font(bold=True, color="7F6000"), stopIfTrue=True))
rg.conditional_formatting.add(area, FormulaRule(formula=[f'LEFT(E{R0},2)="OK"'], fill=GREEN_CF,
                                                font=Font(color="006100"), stopIfTrue=True))
rg.conditional_formatting.add(area, FormulaRule(formula=[f'OR(LEFT(E{R0},1)="—",E{R0}="s.o.")'], fill=GREY_CF,
                                                font=Font(color="808080"), stopIfTrue=True))

# résumé en haut
rg["E4"] = "Alertes (rouge)"
rg["F4"] = f'=COUNTIF({area},"ALERTE*")'
rg["E5"] = "À faire (jaune)"
rg["F5"] = f'=COUNTIF({area},"À FAIRE*")'
rg["E6"] = "OK (vert)"
rg["F6"] = f'=COUNTIF({area},"OK*")'
for rr, fill in ((4, RED_CF), (5, YELLOW_CF), (6, GREEN_CF)):
    rg.cell(rr, 5).font = BOLD
    for cc in (5, 6):
        rg.cell(rr, cc).fill = PatternFill("solid", fgColor=fill.fgColor.rgb)
        rg.cell(rr, cc).border = BOX
rg.cell(R_LAST + 2, 1, ("Note : la règle Max. conversions reste « À FAIRE » tant que le seuil est atteint ; "
                        "ignorez-la si la campagne y est déjà. Les règles d'arrêt sont aussi au plan (CPA +25 % "
                        "après une hausse de CPC : revenir en arrière).")).font = NOTE
widths(rg, [40, 40, 40, 12] + [24] * NC + [12] * 3)
for j in range(len(IND_COLS)):
    if j >= 3 + NC:
        rg.column_dimensions[L(2 + j)].width = 13
rg.freeze_panes = "B4"

# =================================================================== MODE D'EMPLOI
md = ws_mode
md["A1"] = "Mode d'emploi — chaque lundi, 15 minutes"
md["A1"].font = TITLE
hdr(md, 3, ["Étape", "Quoi faire", "Où trouver le chiffre"], height=24)
steps = [
    ("Avant la 1re saisie",
     "Effacer le contenu des 2 lignes orange « EXEMPLE — à effacer » (semaine du 2026-10-12) : sélectionner les "
     "cellules jaunes et la remarque, puis touche Suppr. Ne pas supprimer les lignes.",
     "Feuille « Saisie hebdo », lignes 5 et 6."),
    ("1. Google Ads — période",
     "Ouvrir Campagnes et choisir la période « Semaine dernière (lun. au dim.) ». Repérer les lignes de cette "
     "semaine (colonne Semaine = lundi) dans « Saisie hebdo ».",
     "Google Ads › Campagnes › sélecteur de dates en haut à droite."),
    ("2. Coût, impressions, clics, conversions",
     "Pour chaque campagne active, recopier Coût, Impr., Clics et Conversions (formulaire + appels, actions "
     "principales seulement).",
     "Google Ads › Campagnes, colonnes par défaut. Vérifier les actions dans Objectifs › Conversions."),
    ("3. Parts d'impressions perdues",
     "Recopier « Part d'impressions perdue sur le Réseau de Recherche (budget) » et « (classement) » en % "
     "(taper 12 % ou 0,12). Laisser vide pour Demand Gen.",
     "Colonnes › Modifier les colonnes › Statistiques sur la concurrence."),
    ("4. Notion — leads",
     "Filtrer la base Leads : Source = Google Ads, date de création dans la semaine, regroupé par Campagne. "
     "Leads reçus = toutes les fiches ; Leads qualifiés = Score A, B ou C (D exclu).",
     "Notion › base Leads › champs Source, Campagne, Score. Corporatif → Corporatif & Fêtes ; Mariage/Produits → "
     "Événements privés ; Marque → Marque."),
    ("5. Notion — soumissions et contrats",
     "Soumissions envoyées = fiches passées à « Soumission envoyée » cette semaine. Contrats gagnés = « Dépôt "
     "payé » cette semaine ; Valeur = somme du champ Montant du contrat. On compte la semaine du dépôt, pas celle du lead.",
     "Notion › champs État et Montant du contrat."),
    ("6. Lire la feuille Règles",
     "Traiter d'abord les cellules rouges (ALERTE), puis les jaunes (À FAIRE). Appliquer le changement dans "
     "Google Ads et le noter dans la colonne Remarque de la saisie.",
     "Feuille « Règles », section 2. Le compteur en haut indique le nombre d'alertes."),
    ("7. Vendredi — import des leads qualifiés",
     "Exporter les leads A, B et C avec GCLID et les importer comme « Lead qualifié » ; les dépôts payés comme "
     "« Réservation » avec le montant (sans ça, Google n'apprend pas).",
     "Google Ads › Objectifs › Conversions › Importations. Voir SCRIPTS-LEADS.md, section 8."),
    ("8. Fin de mois et jour 30",
     "Consulter « Synthèse mensuelle » (coût par lead qualifié, par contrat, ROAS). Mettre à jour « Cibles » quand "
     "la marge et le panier réel sont connus ; les règles se recalculent seules.",
     "Feuilles « Synthèse mensuelle » et « Cibles »."),
]
for k, (a, b, c) in enumerate(steps):
    rr = 4 + k
    for cc, v in enumerate((a, b, c), 1):
        cell = md.cell(rr, cc, v)
        cell.alignment, cell.border = WRAP, BOX
    md.cell(rr, 1).font = BOLD
    md.row_dimensions[rr].height = 62
md.cell(4, 1).fill = ORANGE
md.cell(14, 1, "Rappels").font = BOLD
md.cell(15, 1, "• Demand Gen, Microsoft Ads, S-EN et Québec-Lévis : laisser vide tant que la campagne est en pause.")
md.cell(16, 1, "• Un 0 saisi compte comme une donnée ; une cellule vide signifie « pas encore saisi ».")
md.cell(17, 1, "• Pour ajouter une campagne : l'ajouter au tableau Campagnes de « Cibles », puis copier des lignes dans « Saisie hebdo ».")
for rr in (15, 16, 17):
    md.cell(rr, 1).font = NOTE
widths(md, [30, 80, 60])
md.freeze_panes = "A4"

# ordre et finition
wb._sheets = [ws_saisie, ws_synth, ws_reg, cb, ws_mode]
wb.active = 0
for s in wb.worksheets:
    s.sheet_view.zoomScale = 90
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"OK : {OUT} ({LAST_DATA - 4} lignes de saisie, {len(RULES)} règles)")
